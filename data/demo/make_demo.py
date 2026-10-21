"""T-017a: шесть выдуманных документов; ключ ответов не зависит от детектора.

Запуск в api: python /data_rw/demo/make_demo.py --out /data_rw/demo
Проверки: python /data_rw/demo/make_demo.py --self-test
"""

from __future__ import annotations

import argparse
import logging
import sys
import tempfile
import unittest
from html import escape
from pathlib import Path
from random import Random

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

# При запуске файла из bind mount Python не добавляет WORKDIR /app в sys.path.
sys.path.insert(0, str(Path("/app") if Path("/app/app").is_dir()
                       else Path(__file__).resolve().parents[2] / "backend"))
from app.evaluation.generator.values import inn_person, snils  # noqa: E402
from app.services.detection import validators  # noqa: E402

LOG = logging.getLogger(__name__)
NOTICE = "Учебный документ. Все организации, люди и реквизиты вымышлены."
Row = tuple[str, str]
Pages = tuple[tuple[Row, ...], ...]


def scenarios() -> dict[str, Pages]:
    """Содержимое фиксировано; случайность используется только для контрольных сумм."""
    rng = Random(1701)
    snils_value, inn_value = snils(rng), inn_person(rng)
    return {
        "01_tz_publication.docx": ((
            ("title", "Техническое задание на поставку рабочих станций"),
            ("body", "Заказчик: ООО «Пример Закупки». Назначение: для публикации."),
            ("heading", "Предмет и объём поставки"),
            ("body", "Требуется поставить 20 рабочих станций для оборудованного учебного класса. "
             "Каждый комплект включает системный блок, монитор, клавиатуру и мышь. "
             "Оборудование должно быть новым, совместимым с действующей локальной сетью."),
            ("heading", "Требования к исполнению"),
            ("body", "Поставщик доставляет оборудование одной партией в течение "
             "30 календарных дней "
             "после заключения контракта. При приёмке заказчик проверяет комплектность, "
             "работоспособность и наличие эксплуатационных документов на русском языке."),
            ("body", "Гарантийное обслуживание предоставляется в течение 24 месяцев. "
             "Замена неисправного оборудования производится в течение пяти рабочих дней."),
            ("heading", "Контакт по техническим вопросам"),
            ("body", "Контактное лицо: Лебедев Артём Павлович. "
             "Личный телефон: +7 (999) 170-01-01."),
            ("heading", "Расчёт поставщика"),
            ("body", "Себестоимость единицы составляет 48 300 руб. "
             "Показатель включён в приложение к техническому заданию для расчёта предложения."),
            ("body", "Разъяснения по требованиям принимаются через электронную площадку. "
             "Поставщик прикладывает описание предлагаемой комплектации и график доставки."),
        ),),
        "02_contract_clean.docx": ((
            ("title", "Проект контракта на сопровождение закупки"),
            ("body", "Конфиденциально. ООО «Пример Закупки». Назначение: договор с контрагентом."),
            ("heading", "Предмет договора"),
            ("body", "Заказчик ООО «Пример Закупки» поручает исполнителю ООО «Пример Сервис» "
             "подготовку комплекта приёмочных документов и сопровождение поставки рабочих станций. "
             "Срок оказания услуг составляет 30 календарных дней с даты заключения договора."),
            ("heading", "Условия поручения обработки персональных данных"),
            ("body", "Заказчик поручает обработку персональных данных исполнителю в соответствии "
             "с ч. 3 ст. 6 152-ФЗ. Цель обработки: проверка полномочий подписанта и оформление "
             "приёмки. Перечень данных: ФИО и паспорт подписанта. Действия: получение, запись, "
             "систематизация, хранение, использование для приёмки и уничтожение."),
            ("body", "Исполнитель обязан соблюдать конфиденциальность персональных данных, "
             "выполнять требования ст. 19 152-ФЗ и предоставлять заказчику "
             "подтверждение принятых мер. Передача другим обработчикам допускается "
             "только по письменному разрешению заказчика."),
            ("heading", "Меры защиты"),
            ("body", "Доступ предоставляется по именным учётным записям только юристу "
             "и специалисту по приёмке. Применяются шифрование хранилища, защищённый "
             "канал передачи, журналирование "
             "доступа без значений персональных данных, ежедневное резервное копирование. "
             "Бумажный экземпляр находится в запираемом шкафу. Об инциденте исполнитель сообщает "
             "заказчику в течение 24 часов."),
        ), (
            ("heading", "Согласие подписанта"),
            ("body", "Подписант: Кузнецов Глеб Романович. Паспорт: 99 17 170202."),
            ("body", "Подписант даёт согласие на обработку персональных данных заказчиком "
             "и исполнителем для проверки полномочий и оформления приёмки. Согласие охватывает "
             "ФИО и паспорт, действия и меры защиты, перечисленные в договоре. "
             "Согласие подписанта приложено к договору; отзыв направляется заказчику письменно."),
            ("heading", "Срок хранения и уничтожение"),
            ("body", "Срок хранения персональных данных: 3 года после завершения исполнения "
             "договора. По истечении срока исполнитель удаляет электронные копии и уничтожает "
             "бумажные экземпляры с составлением акта в течение 10 рабочих дней."),
            ("heading", "Лист согласования"),
            ("body", "Юридический отдел: согласовано. Служба информационной безопасности: "
             "согласовано. Отдел закупок: согласовано. Визы подразделений приложены "
             "к экземпляру договора. Дата согласования: 20.10.2026."),
            ("body", "Услуги принимаются по акту после проверки комплектности документов. "
             "При замечаниях исполнитель исправляет комплект в течение трёх рабочих дней. "
             "Обмен документами осуществляется через согласованный защищённый канал."),
        )),
        "03_application_pd.pdf": ((
            ("title", "Заявка участника на поставку оборудования"),
            ("body", "Участник: ООО «Пример Поставка». Заказчик: ООО «Пример Закупки». "
             "Назначение: договор с контрагентом."),
            ("heading", "Предложение участника"),
            ("body", "Просим рассмотреть предложение на поставку 20 рабочих станций. "
             "Подтверждаем готовность доставить новые комплекты в течение 30 календарных дней. "
             "Условия гарантии и приёмки принимаем в соответствии с техническим заданием."),
            ("heading", "Сведения о представителе"),
            ("body", "Представитель участника: Соколов Роман Денисович."),
            ("body", "Паспорт: 99 17 170303."),
            ("body", f"СНИЛС: {snils_value}."),
            ("body", f"ИНН физического лица: {inn_value}."),
            ("body", "Дата рождения: 17.03.1987."),
            ("heading", "Порядок рассмотрения"),
            ("body", "Представитель уполномочен обсуждать комплектность поставки и передавать "
             "приёмочные документы. Просим включить перечисленные сведения в проект контракта. "
             "Техническое предложение и график доставки прилагаются отдельными файлами."),
            ("body", "Приёмка проводится после проверки комплектности каждой рабочей станции. "
             "Участник обеспечивает упаковку, разгрузку и первоначальную проверку оборудования."),
        ),),
        "05_anketa_special.docx": ((
            ("title", "Анкета работника для организации приёмки"),
            ("body", "ООО «Пример Закупки». Назначение: внутренний документ."),
            ("body", "Анкета передаётся в отдел закупок при формировании группы приёмки "
             "оборудования. Сведения внесены кадровой службой для планирования участия работника."),
            ("heading", "Сведения о работнике"),
            ("body", "Работник: Морозов Тимур Олегович. Должность: специалист по приёмке."),
            ("body", "Состояние здоровья: инвалид II группы."),
            ("body", "Общественная деятельность: член профсоюза."),
            ("heading", "Участие в закупке"),
            ("body", "Работнику предлагается проверить комплектность поставки и наличие "
             "эксплуатационных документов. Выезд на склад планируется в рабочее время. "
             "Результаты проверки передаются руководителю группы приёмки."),
            ("body", "Рабочее место оборудуется до начала поставки. График участия уточняется "
             "после получения уведомления о готовности оборудования к отгрузке. "
             "Отдел закупок использует анкету при распределении обязанностей в группе."),
        ),),
        "06_medical_tz.docx": ((
            ("title", "Техническое задание на поставку лекарственных препаратов"),
            ("body", "Заказчик: ООО «Пример Клиника». Назначение: для публикации."),
            ("heading", "Предмет закупки"),
            ("body", "Требуется поставка зарегистрированных лекарственных препаратов для "
             "пополнения складского запаса. Номенклатура включает препараты для лечения "
             "онкологических заболеваний и средства для пациентов с диабетом. "
             "Количество и формы выпуска указываются в спецификации к контракту."),
            ("heading", "Требования к товару"),
            ("body", "Препараты поставляются в заводской упаковке с инструкцией на русском языке. "
             "Остаточный срок годности на дату поставки должен составлять не менее 12 месяцев. "
             "Поставщик соблюдает условия перевозки и температурный режим производителя."),
            ("heading", "Поставка и приёмка"),
            ("body", "Поставка осуществляется партиями по заявке склада в течение 10 рабочих дней. "
             "Приёмка включает проверку количества, целостности упаковки, маркировки и документов "
             "о качестве. Серии и сроки годности отражаются в товаросопроводительных документах."),
            ("heading", "Общие контакты организации"),
            ("body", "Информационная линия организации: 8-800-999-17-06. "
             "Общий почтовый ящик: info@example.ru. Сайт: https://example.org."),
            ("body", "Вопросы по закупке направляются через электронную площадку. "
             "Техническое задание описывает требования к ассортименту и логистике поставки."),
        ),),
    }


def write_docx(path: Path, pages: Pages) -> None:
    """Обычный закупочный документ A4, без колонтитулов с дополнительными находками."""
    if not pages or any(not page for page in pages):
        raise ValueError("Документ и его страницы не должны быть пустыми")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name, size in (("Normal", 11), ("Title", 19), ("Heading 1", 13)):
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in list(fonts.attrib):
            if attribute.endswith("Theme"):
                del fonts.attrib[attribute]
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
    document.styles["Normal"].paragraph_format.line_spacing = 1.1
    document.core_properties.author = "Демо T-017a"
    document.core_properties.last_modified_by = "Демо T-017a"
    for page_index, rows in enumerate(pages):
        if page_index:
            document.add_page_break()
        for kind, text in rows:
            document.add_paragraph(text, {"title": "Title", "heading": "Heading 1"}.get(kind))
    document.add_paragraph(NOTICE)
    document.save(path)


def write_pdf(path: Path, pages: Pages) -> None:
    """MuPDF встраивает кириллицу; переполнение не уменьшает шрифт незаметно."""
    if not pages or any(not page for page in pages):
        raise ValueError("Документ и его страницы не должны быть пустыми")
    css = "body {font-family:sans-serif;font-size:11pt;line-height:1.35;} "
    css += "h1 {font-size:19pt;} h2 {font-size:13pt;} p {margin:0 0 10pt;}"
    with pymupdf.open() as document:
        for rows in pages:
            page = document.new_page(width=595, height=842)
            tags = {"title": "h1", "heading": "h2", "body": "p"}
            html = "".join(f"<{tags[kind]}>{escape(text)}</{tags[kind]}>" for kind, text in rows)
            spare, _ = page.insert_htmlbox(
                pymupdf.Rect(55, 45, 540, 795), html + f"<p>{NOTICE}</p>",
                css=css, scale_low=1,
            )
            if spare < 0:
                raise ValueError("Текст не помещается на страницу PDF")
        document.set_metadata({"title": "Демонстрационная заявка", "author": "Демо T-017a"})
        document.save(path)


def write_estimate(path: Path) -> None:
    """Стоимость и наценка находятся в обычных ячейках сметы, без подгонки под детектор."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Смета"
    rows = [
        ["Смета поставки рабочих станций"],
        ["ООО «Пример Поставка». Назначение: внутренний документ."],
        ["Расчёт предложения на 20 комплектов для учебного класса."],
        ["Показатель", "Значение", "Единица"],
        ["Количество комплектов", 20, "шт."],
        ["Себестоимость единицы", 48300, "руб."],
        ["Наценка", 0.25, "%"],
        ["Цена единицы", 60375, "руб."],
        ["Сумма предложения", 1207500, "руб."],
        ["Расчёт включает доставку и первоначальную проверку комплектности."],
        ["Поставка одной партией в течение 30 календарных дней."],
        ["Расчёт подготовлен для обсуждения предложения отделом закупок."],
        [NOTICE],
    ]
    for row in rows:
        sheet.append(row)
    for row_index in (1, 2, 3, 10, 11, 12, 13):
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=3)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=11, color="000000")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=4, max_row=9):
        for cell in row:
            side = Side(style="thin", color="D9D9D9")
            cell.border = Border(left=side, right=side, top=side, bottom=side)
    for cell in sheet[4]:
        cell.fill, cell.font = PatternFill("solid", fgColor="DCE6F1"), Font(bold=True, size=11)
    sheet["A1"].font = Font(name="Arial", size=18, bold=True)
    sheet.column_dimensions["A"].width = 43
    sheet.column_dimensions["B"].width = 19
    sheet.column_dimensions["C"].width = 12
    for index in range(1, 14):
        sheet.row_dimensions[index].height = 34 if index in (2, 3, 10, 12, 13) else 29
    for index in (6, 8, 9):
        sheet.cell(index, 2).number_format = '#,##0.00'
    sheet["B7"].number_format = "0%"
    sheet.freeze_panes = "B5"
    sheet.print_options.horizontalCentered = True
    sheet.print_area = "A1:C13"
    sheet.page_setup.orientation, sheet.page_setup.paperSize = "portrait", sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    workbook.save(path)
    workbook.close()


def generate(out: Path) -> None:
    """Перезаписывает только шесть известных имён, не меняет answers.yaml и README."""
    if out.exists() and not out.is_dir():
        raise ValueError("--out должен указывать на каталог")
    out.mkdir(parents=True, exist_ok=True)
    for filename, pages in scenarios().items():
        writer = write_pdf if filename.endswith(".pdf") else write_docx
        writer(out / filename, pages)
    write_estimate(out / "04_estimate.xlsx")
    LOG.info("T-017a: создано 6 синтетических документов")


class DemoTests(unittest.TestCase):
    """Проверки генерации без вызова детектора, правил, сети и LLM."""

    def test_values_and_reproducibility(self) -> None:
        self.assertEqual(scenarios(), scenarios())
        for seed in (0, 1701):
            rng = Random(seed)
            self.assertTrue(validators.is_valid_snils(snils(rng)))
            self.assertTrue(validators.is_valid_inn(inn_person(rng)))

    def test_docx_success_and_empty(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.docx"
            write_docx(path, ((('title', 'Тест'), ('body', 'Проверка')),))
            self.assertIn("Проверка", [p.text for p in Document(path).paragraphs])
            with self.assertRaises(ValueError):
                write_docx(path, ())

    def test_pdf_success_and_overflow(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.pdf"
            write_pdf(path, scenarios()["03_application_pd.pdf"])
            with pymupdf.open(path) as document:
                self.assertEqual(len(document), 1)
                self.assertIn("СНИЛС", document[0].get_text())
            with self.assertRaises(ValueError):
                write_pdf(path, ((('body', 'Длинный текст ' * 5000),),))
            with self.assertRaises(ValueError):
                write_pdf(path, ())

    def test_estimate_values_and_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.xlsx"
            write_estimate(path)
            workbook = load_workbook(path, data_only=True)
            sheet = workbook["Смета"]
            self.assertEqual(sheet["B8"].value, sheet["B6"].value * (1 + sheet["B7"].value))
            self.assertEqual(sheet["B9"].value, sheet["B5"].value * sheet["B8"].value)
            self.assertEqual(sheet["B7"].number_format, "0%")
            self.assertEqual(len(workbook.sheetnames), 1)
            workbook.close()

    def test_generate_preserves_other_files_and_repeat(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            key = out / "answers.yaml"
            key.write_text("unchanged", encoding="utf-8")
            generate(out)
            generate(out)
            self.assertEqual(len(list(out.iterdir())), 7)
            self.assertEqual(key.read_text(encoding="utf-8"), "unchanged")

    def test_generate_rejects_file_as_output(self) -> None:
        with tempfile.NamedTemporaryFile() as target:
            with self.assertRaises(ValueError):
                generate(Path(target.name))

    def test_medical_traps_have_no_person(self) -> None:
        text = " ".join(text for page in scenarios()["06_medical_tz.docx"] for _, text in page)
        self.assertIn("info@example.ru", text)
        self.assertIn("8-800-999-17-06", text)
        self.assertIn("онкологических заболеваний", text)
        self.assertNotIn("СНИЛС", text)
        self.assertNotIn("Работник:", text)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path(__file__).parent)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.self_test:
        suite = unittest.defaultTestLoader.loadTestsFromTestCase(DemoTests)
        result = unittest.TextTestRunner(verbosity=2).run(suite)
        raise SystemExit(0 if result.wasSuccessful() else 1)
    generate(args.out)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    main()
