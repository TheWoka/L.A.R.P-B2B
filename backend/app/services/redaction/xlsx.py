"""Обезличивание .xlsx на openpyxl: замена значений ПДн в ячейках. Задача T-026.

Внутренний модуль пакета redaction — не контракт.
"""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment

from app.services.redaction.values import PdValue, replace_values, value_pattern

_PROPERTY_FIELDS = ("title", "subject", "keywords", "description", "category")
_AUTHOR_FIELDS = ("creator", "lastModifiedBy")


def redact_xlsx(source: Path, values: list[PdValue]) -> bytes:
    """Возвращает байты копии .xlsx, где значения ПДн заменены метками.

    Строковые ячейки и формулы — замена подстрок; число или дата, в текстовом виде
    которых найдено значение ПДн, заменяются строкой-меткой целиком. Примечания
    к ячейкам — та же замена. Формулы пересчитываются при открытии файла.
    """
    with source.open("rb") as stream:
        workbook = load_workbook(stream)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                _redact_cell(cell, values)
    properties = workbook.properties
    for name in _AUTHOR_FIELDS:
        setattr(properties, name, None)
    for name in _PROPERTY_FIELDS:
        if current := getattr(properties, name):
            setattr(properties, name, replace_values(current, values))
    # Значения формул openpyxl не хранит: Excel и LibreOffice пересчитают их при открытии.
    workbook.calculation.fullCalcOnLoad = True
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _redact_cell(cell: Cell, values: list[PdValue]) -> None:
    """Меняет значение и примечание ячейки, если в них есть ПДн."""
    if cell.comment is not None:
        text = replace_values(cell.comment.text, values)
        if text != cell.comment.text:
            cell.comment = Comment(text, "")
    if cell.value is None:
        return
    if isinstance(cell.value, str):
        cell.value = replace_values(cell.value, values)
        return
    # extraction/xlsx.py видит число как str(value): «729423480941», «1987-03-17 00:00:00».
    as_text = str(cell.value)
    for value in values:
        if value_pattern(value).search(as_text):
            cell.value = value.label
            return
