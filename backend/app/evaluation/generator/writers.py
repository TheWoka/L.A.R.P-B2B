"""Запись блоков документов имеющимися библиотеками, без новых зависимостей."""

from html import escape
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook

from app.evaluation.generator.templates import Block


def write_docx(path: Path, blocks: tuple[Block, ...]) -> None:
    document = Document()
    for block in blocks:
        if block.kind == "table":
            width = max(len(row) for row in block.rows)
            table = document.add_table(rows=0, cols=width)
            for row in block.rows:
                for cell, text in zip(table.add_row().cells, row, strict=False):
                    cell.text = text
        else:
            for row in block.rows:
                document.add_paragraph(" | ".join(row))
    document.save(str(path))


def write_xlsx(path: Path, blocks: tuple[Block, ...]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Смета"
    for block in blocks:
        for row in block.rows:
            sheet.append(list(row))
    workbook.save(path)
    workbook.close()


def write_txt(path: Path, blocks: tuple[Block, ...]) -> None:
    text = "\n".join(" | ".join(row) for block in blocks for row in block.rows)
    path.write_text(text + "\n", encoding="utf-8")


def write_pdf(path: Path, blocks: tuple[Block, ...]) -> None:
    """Широкие страницы исключают переносы внутри размечаемых фраз."""
    with pymupdf.open() as document:
        page = document.new_page(width=1190, height=842)
        top = 30.0
        for block in blocks:
            for row in block.rows:
                if top + 70 > 812:
                    page = document.new_page(width=1190, height=842)
                    top = 30.0
                html = "<p>" + escape(" | ".join(row)) + "</p>"
                spare, _ = page.insert_htmlbox(
                    pymupdf.Rect(30, top, 1160, top + 70), html,
                    css="* {font-family: sans-serif; font-size: 11pt;}", scale_low=1,
                )
                if spare < 0:
                    raise ValueError("PDF: блок не помещается на страницу")
                top += 70 - spare + 8
        document.save(path)


WRITERS = {"docx": write_docx, "xlsx": write_xlsx, "txt": write_txt, "pdf": write_pdf}
