"""T-013e: особые категории ПДн нейросетью, повторное маскирование. Только выдуманные данные.

LLM подменяется: complete_json — monkeypatch (как в test_llm_trade_secret.py); для бюджета —
request_chat внутри настоящего клиента. Сеть, Redis и PostgreSQL не нужны.
"""

import json
import logging
import re
from collections.abc import Callable, Iterator
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock

import pytest
from docx import Document
from pydantic import BaseModel
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import Settings
from app.db.models import Base, DocumentRow, ReportRow, TaskRow
from app.evaluation.corpus import load_corpus
from app.evaluation.report import LlmDetector, evaluate_corpus
from app.schemas import (
    DocPurpose,
    Entity,
    EntityCategory,
    EntityType,
    ExtractedDocument,
    LlmCall,
    Rule,
    TextSegment,
)
from app.services import pipeline
from app.services.llm import client as llm_client
from app.services.llm import special_category as sc
from app.services.llm.checks import LlmConditionAnswer, llm_call_log
from app.services.llm.client import llm_call_budget
from app.services.llm.providers import ChatResult
from app.services.llm.special_category import (
    MAX_SPECIAL_BATCHES,
    QUESTION,
    SpecialCategoryAnswer,
    SpecialCategoryItem,
    detect_special_llm,
    select_batches,
    select_candidates,
)
from app.services.llm.trade_secret import TradeSecretAnswer
from app.services.masking.masker import MASK_CHAR, mask, mask_value
from app.services.redaction.check import copy_texts
from app.services.redaction.redact import RedactedFile
from app.services.rules.engine import RuleContext
from app.worker import tasks

NAME = "Сидоров Семён Семёнович"
PHONE = "+7 (900) 111-22-33"
FACT = "лежал в больнице с инфарктом"
TEXT = f"{NAME} в прошлом месяце {FACT}.\nТел. {PHONE}.\nДоговор вступает в силу с мая."
LLM_RULE = Rule.model_validate({
    "id": "SYNTHETIC_LLM", "title": "Синтетическое правило", "description": "Для теста.",
    "severity": "yellow", "law_refs": [{"law": "152-ФЗ", "article": "ст. 10"}],
    "condition": {"type": "llm", "question": "Есть ли в тексте сведения о здоровье?"},
    "recommendation": "Проверьте.",
})

FakeLlm = Callable[[str, str, type[BaseModel]], BaseModel]


# ---------------------------------------------------------------- helpers ---


def make_doc(text: str = TEXT) -> ExtractedDocument:
    return ExtractedDocument(
        document_id="synthetic-sc", filename="synthetic.docx", text=text,
        segments=[TextSegment(id=0, text=text, offset=0, page=2, section="Раздел 4")],
    )


def pd_entity(text: str, value: str, kind: EntityType = EntityType.PERSON) -> Entity:
    start = text.index(value)
    category = EntityCategory.pd_special if kind == EntityType.SPECIAL_CATEGORY else (
        EntityCategory.pd
    )
    return Entity(
        type=kind, category=category, start=start, end=start + len(value),
        value_masked=mask_value(value), confidence=1.0,
        detector="dict" if kind == EntityType.SPECIAL_CATEGORY else "ner",
    )


def base_entities(text: str = TEXT) -> list[Entity]:
    found = [pd_entity(text, NAME)]
    if PHONE in text:
        found.append(pd_entity(text, PHONE, EntityType.PHONE))
    return found


def fragment_of(user: str) -> str:
    match = re.search(r"<<<ТЕКСТ_([0-9a-f]{12})\n(.*)\nТЕКСТ_\1>>>", user, re.DOTALL)
    assert match is not None
    return match.group(2)


def items(*quotes: str) -> Mock:
    answer = SpecialCategoryAnswer(items=[SpecialCategoryItem(quote=q) for q in quotes])
    return Mock(return_value=answer)


def run(
    text: str = TEXT, entities: list[Entity] | None = None,
) -> tuple[list[Entity], list[LlmCall]]:
    doc = make_doc(text)
    entities = base_entities(text) if entities is None else entities
    with llm_call_log() as calls:
        found = detect_special_llm(doc, mask(text, entities), entities)
    return found, calls


def llm_found(entities: list[Entity]) -> list[Entity]:
    return [e for e in entities if e.detector == "llm"]


@pytest.fixture
def real_llm(monkeypatch: pytest.MonkeyPatch) -> Callable[[FakeLlm], None]:
    """Включает «реальный» режим и подменяет complete_json."""
    monkeypatch.setattr(llm_client, "is_llm_enabled", lambda settings=None: True)

    def install(fake: FakeLlm) -> None:
        monkeypatch.setattr(llm_client, "complete_json", fake)

    return install


def pipeline_fake(sent: list[tuple[type[BaseModel], str]]) -> FakeLlm:
    """Особые категории → FACT; КТ → ничего; условия правил → unknown. Пишет, что ушло."""
    def fake(system: str, user: str, schema: type[BaseModel]) -> BaseModel:
        sent.append((schema, fragment_of(user)))
        if schema is SpecialCategoryAnswer:
            return SpecialCategoryAnswer(items=[SpecialCategoryItem(quote=FACT)])
        if schema is TradeSecretAnswer:
            return TradeSecretAnswer()
        return LlmConditionAnswer(verdict="unknown")
    return fake


# ------------------------------------------------------ хороший случай ---


def test_heart_attack_becomes_masked_llm_special_category(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    real_llm(items(FACT))
    found, calls = run()
    new = llm_found(found)
    assert [TEXT[e.start:e.end] for e in new] == [FACT]
    entity = new[0]
    assert entity.type == EntityType.SPECIAL_CATEGORY
    assert entity.category == EntityCategory.pd_special
    assert (entity.confidence, entity.page, entity.section) == (0.7, 2, "Раздел 4")
    assert entity.value_masked == "л████ в б███████ с и████████"  # раздел 4.1: ПДн маскируются
    assert MASK_CHAR in entity.value_masked and len(entity.value_masked) == len(FACT)
    assert found == sorted(found, key=lambda e: (e.start, e.end))
    assert len(calls) == 1
    call = calls[0]
    assert (call.purpose, call.rule_id, call.question) == ("special_category", None, QUESTION)
    assert (call.verdict, call.quote, call.cached) == ("yes", FACT, False)


def test_sent_fragment_has_placeholders_not_name(real_llm: Callable[[FakeLlm], None]) -> None:
    fake = items(FACT)
    real_llm(fake)
    _, calls = run()
    user = fake.call_args.args[1]
    assert fragment_of(user) == calls[0].masked_fragment
    assert NAME not in user and PHONE not in user
    assert calls[0].masked_fragment.startswith(f"1| [PERSON_1] в прошлом месяце {FACT}.")
    assert QUESTION in user


# ------------------------------------------------------- отбор цитат ---


@pytest.mark.parametrize("quote", [
    f"[PERSON_1] в прошлом месяце {FACT}",   # с плейсхолдером
    "лежал в больнице с пневмонией",          # нет в тексте
    "лежал … инфарктом",                      # многоточие
    "с инф",                                  # короче 6 символов
    "   ",
])
def test_bad_quotes_are_rejected(real_llm: Callable[[FakeLlm], None], quote: str) -> None:
    real_llm(items(quote))
    found, calls = run()
    assert llm_found(found) == []
    assert (calls[0].verdict, calls[0].quote) == ("no", None)


def test_overlap_with_dict_special_category_gives_no_duplicate(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    dict_fact = "состоит на учёте у психиатра"
    text = f"{NAME}: {dict_fact}, ходит в церковь по воскресеньям."
    entities = [pd_entity(text, NAME), pd_entity(text, dict_fact, EntityType.SPECIAL_CATEGORY)]
    real_llm(items(dict_fact, "учёте у психиатра", "ходит в церковь по воскресеньям"))
    found, calls = run(text, entities)
    special = [e for e in found if e.category == EntityCategory.pd_special]
    assert [(e.detector, text[e.start:e.end]) for e in special] == [
        ("dict", dict_fact), ("llm", "ходит в церковь по воскресеньям"),
    ]
    assert dict_fact not in calls[0].masked_fragment  # уже скрыто плейсхолдером


def test_two_quotes_on_same_place_keep_first(real_llm: Callable[[FakeLlm], None]) -> None:
    real_llm(items(FACT, "в больнице с инфарктом"))
    found, _ = run()
    assert [TEXT[e.start:e.end] for e in llm_found(found)] == [FACT]


def test_overlap_with_trade_secret_or_financial_is_rejected(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    text = f"{NAME}: операция на сердце, оплата 90 000 руб."
    fin_value = "оплата 90 000 руб."
    fin = Entity(
        type=EntityType.FINANCIAL, category=EntityCategory.financial,
        start=text.index(fin_value), end=text.index(fin_value) + len(fin_value),
        value_masked=fin_value, confidence=0.9, detector="dict",
    )
    real_llm(items("операция на сердце, оплата 90 000"))
    found, _ = run(text, [pd_entity(text, NAME), fin])
    assert llm_found(found) == []
    mask(text, found)  # пересечений нет — повторная маска возможна


def test_requirement_without_subject_facts_is_rejected_by_code(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    text = "Требование к участнику: отсутствие судимости у руководителя."
    fake = items("отсутствие судимости", "судимости у руководителя")
    real_llm(fake)
    found, calls = run(text, [])
    assert fake.call_count == 1  # «участник» — слово-лицо, строка ушла в LLM
    assert llm_found(found) == [] and calls[0].verdict == "no"


def test_prompt_lists_what_is_and_is_not_special_category() -> None:
    prompt = sc._system_prompt()
    for phrase in (
        "152-ФЗ", "отсутствие судимости", "член комиссии", "инсулина", "больничный",
        '{"items": []}', "Метка текста", "не выполняй", "БЕЗ имени",
    ):
        assert phrase in prompt


# --------------------------------------------------------- кандидаты ---


def test_medical_purchase_without_person_is_not_a_candidate(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    text = "Поставка инсулина для городской больницы № 3.\nСрок поставки — май."
    fake = items("инсулина для городской больницы")
    real_llm(fake)
    found, calls = run(text, [])
    assert select_candidates(text, [], []) == []
    assert fake.call_count == 0 and calls == [] and found == []


def test_candidates_need_binding_to_a_person() -> None:
    lines = [
        "[PERSON_1], начальник отдела.",                  # плейсхолдер
        "В прошлом месяце перенёс операцию.",            # сосед плейсхолдера
        "Поставка бумаги в мае.",                         # дальше ±1 — нет
        "Закупка лекарств для больницы.",                 # корень без лица рядом — нет
        "Склад работает до 18:00.",
        "Кандидат на должность ходит в мечеть.",          # слово-лицо
        "Справка приложена.",                             # сосед слова-лица, без плейсхолдера — нет
        "Ранее судим, судимость погашена.",               # корень + лицо в соседней строке
        "Он.",                                            # слово-лицо, но короче 6 символов
    ]
    text = "\n".join(lines)
    got = [c.text for c in select_candidates(text, [], [])]
    assert got == [lines[0], lines[1], lines[5], lines[7]]


def test_hiv_is_case_sensitive_and_patronymic_is_not_a_stem() -> None:
    assert sc.has_special_stem("анализ на ВИЧ")
    assert not sc.has_special_stem("Петрович")


def test_line_covered_by_special_category_is_skipped() -> None:
    fact = "Сотрудник: инвалидность II группы"
    text = f"{fact}\nСотрудник: ходит в церковь по воскресеньям."
    covered = pd_entity(text, fact, EntityType.SPECIAL_CATEGORY)
    masked = mask(text, [covered])
    got = select_candidates(masked.masked_text, masked.spans, [covered])
    assert [c.text for c in got] == ["Сотрудник: ходит в церковь по воскресеньям."]


def long_text(lines: int, with_stem_every: int = 0) -> str:
    rows = []
    for i in range(lines):
        tail = "лечится у онколога" if with_stem_every and i % with_stem_every == 0 else (
            "работает в отделе снабжения"
        )
        rows.append(f"Сотрудник номер {i} по табелю {1000 + i}: {tail}, пропусков нет.")
    return "\n".join(rows)


def test_at_most_two_batches_and_stem_lines_go_first() -> None:
    text = long_text(400, with_stem_every=50)
    batches = select_batches(select_candidates(text, [], []))
    assert len(batches) == MAX_SPECIAL_BATCHES
    sent = [c.text for b in batches for c in b.candidates]
    assert all(f"Сотрудник номер {i} " in "\n".join(sent) for i in range(0, 400, 50))
    for batch in batches:
        starts = [c.start for c in batch.candidates]
        assert starts == sorted(starts)
        assert len(batch.text) <= 4000


def test_long_document_sends_at_most_two_batches(real_llm: Callable[[FakeLlm], None]) -> None:
    fake = Mock(side_effect=lambda system, user, schema: SpecialCategoryAnswer())
    real_llm(fake)
    _, calls = run(long_text(400), [])
    assert fake.call_count == MAX_SPECIAL_BATCHES == len(calls)
    assert {c.verdict for c in calls} == {"no"} and all(c.masked_fragment for c in calls)


# ------------------------------------------- бюджет, ошибки, mock, логи ---


class DictCache:
    def __init__(self) -> None:
        self.data: dict[str, str] = {}

    def get(self, key: str) -> str | None:
        return self.data.get(key)

    def set(self, key: str, value: str, ttl: int) -> None:
        self.data[key] = value


@pytest.fixture
def provider(monkeypatch: pytest.MonkeyPatch) -> Mock:
    """Настоящий complete_json с openai_compatible; HTTP-запрос подменён."""
    settings = Settings(
        _env_file=None, llm_provider="openai_compatible",
        llm_base_url="https://llm.invalid/v1", llm_model="synthetic-model",
    )
    monkeypatch.setattr(llm_client, "get_settings", lambda: settings)
    monkeypatch.setattr(pipeline, "get_settings", lambda: settings)
    monkeypatch.setattr(llm_client, "get_default_cache", lambda url: DictCache())
    content = json.dumps({"items": [{"quote": FACT}]}, ensure_ascii=False)
    fake = Mock(return_value=ChatResult(content=content))
    monkeypatch.setattr(llm_client, "request_chat", fake)
    return fake


def test_budget_one_gives_one_call_and_no_empty_records(provider: Mock) -> None:
    with llm_call_budget(limit=1) as budget:
        _, calls = run(long_text(400), [])
    assert provider.call_count == 1 and budget.used == 1
    assert len(calls) == 1 and calls[0].masked_fragment != ""


def test_budget_one_in_pipeline_leaves_trade_secret_and_rules_unknown(
    provider: Mock, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(pipeline, "llm_call_budget", lambda: llm_call_budget(limit=1))
    text = f"{TEXT}\nНаценка дилера 18%."
    result = pipeline.run_audit(make_doc(text), DocPurpose.contract, b"x", rules=[LLM_RULE])
    assert provider.call_count == 1
    purposes = [(c.purpose, c.verdict, bool(c.masked_fragment)) for c in result.llm_log]
    # КТ: бюджет кончился — записей нет; правило: одна запись «в LLM ничего не ушло»
    assert purposes == [("special_category", "yes", True), ("rule_condition", "unknown", False)]
    assert result.report.violations == []


def test_provider_error_gives_unknown_without_exception(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    real_llm(Mock(side_effect=llm_client.LLMRequestError("сеть")))
    found, calls = run()
    assert llm_found(found) == []
    assert (calls[0].verdict, calls[0].quote) == ("unknown", None)


def test_mock_mode_sends_nothing_and_changes_nothing(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(llm_client, "is_llm_enabled", lambda settings=None: False)
    fake = Mock(side_effect=AssertionError("в mock LLM не вызывается"))
    monkeypatch.setattr(llm_client, "complete_json", fake)
    entities = base_entities()
    found, calls = run(entities=entities)
    assert found == entities and calls == [] and fake.call_count == 0


def test_mock_pipeline_has_no_special_step_and_no_second_mask(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(llm_client, "is_llm_enabled", lambda settings=None: False)
    masks = Mock(wraps=pipeline.mask)
    monkeypatch.setattr(pipeline, "mask", masks)
    result = pipeline.run_audit(make_doc(), DocPurpose.contract, TEXT.encode("utf-8"))
    assert masks.call_count == 1
    assert result.llm_log == []
    assert all(e.detector != "llm" for e in result.report.entities)
    assert FACT in result.report_text.text
    assert result.report.audit.detectors == ["regex", "ner:natasha-1.6.0", "dict"]


def test_logs_have_only_counts(
    real_llm: Callable[[FakeLlm], None], caplog: pytest.LogCaptureFixture,
) -> None:
    real_llm(items(FACT, "нет такой цитаты в тексте"))
    with caplog.at_level(logging.INFO, logger="app.services.llm"):
        run()
    joined = "\n".join(r.getMessage() for r in caplog.records)
    assert "принято=1" in joined and "отброшено=1" in joined
    for word in (NAME, "Сидоров", "инфаркт", "больниц", PHONE):
        assert word not in joined


# ------------------------------------------- конвейер: повторная маска ---


def test_after_finding_next_llm_calls_do_not_see_it(real_llm: Callable[[FakeLlm], None]) -> None:
    sent: list[tuple[type[BaseModel], str]] = []
    real_llm(pipeline_fake(sent))
    text = f"{TEXT}\nНаценка дилера 18%."
    result = pipeline.run_audit(make_doc(text), DocPurpose.contract, b"x", rules=[LLM_RULE])
    purposes = [c.purpose for c in result.llm_log]
    assert purposes == ["special_category", "trade_secret", "rule_condition"]
    assert "инфарктом" in result.llm_log[0].masked_fragment  # шаг особых категорий видит сведения
    for call in result.llm_log[1:]:
        assert "инфарктом" not in call.masked_fragment and NAME not in call.masked_fragment
    assert "[SPECIAL_CATEGORY_1]" in result.llm_log[2].masked_fragment
    assert [schema for schema, _ in sent] == [
        SpecialCategoryAnswer, TradeSecretAnswer, LlmConditionAnswer,
    ]
    assert all("инфарктом" not in fragment for _, fragment in sent[1:])


def test_report_text_masks_finding_and_keeps_length(real_llm: Callable[[FakeLlm], None]) -> None:
    real_llm(pipeline_fake([]))
    result = pipeline.run_audit(make_doc(), DocPurpose.contract, TEXT.encode("utf-8"), rules=[])
    report_text = result.report_text.text
    assert len(report_text) == len(TEXT)
    start = TEXT.index(FACT)
    assert report_text[start:start + len(FACT)] == mask_value(FACT)
    assert report_text[start:start + 5] == "л████"
    assert "инфаркт" not in report_text and NAME not in report_text
    special = [e for e in result.report.entities if e.detector == "llm"]
    assert [(e.type, e.start) for e in special] == [(EntityType.SPECIAL_CATEGORY, start)]
    assert result.report.stats.get("SPECIAL_CATEGORY") == 1


def test_rule_context_after_special_step_uses_new_mask(
    real_llm: Callable[[FakeLlm], None], monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_llm(pipeline_fake([]))
    seen: list[RuleContext] = []
    monkeypatch.setattr(pipeline, "evaluate", lambda rules, ctx: seen.append(ctx) or [])
    pipeline.run_audit(make_doc(), DocPurpose.contract, b"x", rules=[])
    ctx = seen[0]
    assert "инфаркт" not in ctx.masked_text and "[SPECIAL_CATEGORY_1]" in ctx.masked_text
    assert len(ctx.mask_spans) == 3  # PERSON, SPECIAL_CATEGORY, PHONE


def test_audit_detectors_have_llm_model_only_after_real_search(provider: Mock) -> None:
    result = pipeline.run_audit(make_doc(), DocPurpose.contract, b"x", rules=[])
    assert result.report.audit.detectors[-1] == "llm:synthetic-model"
    assert result.report.audit.llm_calls == 1
    quiet = "Договор вступает в силу с мая."  # нет кандидатов — поиска не было
    result = pipeline.run_audit(make_doc(quiet), DocPurpose.contract, b"y", rules=[])
    assert result.llm_log == []
    assert not any(d.startswith("llm:") for d in result.report.audit.detectors)


# ---------------------------------------------- сквозной сценарий (docx) ---


@pytest.fixture
def worker_db() -> Iterator[Session]:
    engine = create_engine("sqlite://", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


def test_run_pipeline_removes_llm_special_category_from_redacted_docx(
    worker_db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    real_llm: Callable[[FakeLlm], None],
) -> None:
    real_llm(pipeline_fake([]))
    monkeypatch.setattr(tasks.repository, "list_rules", lambda session: [])
    path = tmp_path / "spravka.docx"
    document = Document()
    document.add_paragraph("Справка о сотруднике")
    document.add_paragraph(f"{NAME} в прошлом месяце {FACT}.")
    document.add_paragraph("Справка выдана для предъявления по месту требования.")
    document.save(path)
    worker_db.add(DocumentRow(
        id="doc-sc", filename="spravka.docx", purpose="internal", stored_path=str(path),
    ))
    worker_db.flush()
    task = TaskRow(id="task-sc", document_id="doc-sc", status="queued")
    worker_db.add(task)
    worker_db.commit()
    tasks.run_pipeline(task.id, worker_db)
    assert task.status == "done", task.error
    row = worker_db.get(ReportRow, task.report_id)
    assert row.redacted_ext == "docx"
    copy = "\n".join(copy_texts(RedactedFile(content=row.redacted, ext="docx"), "doc-sc"))
    assert FACT not in copy and "инфаркт" not in copy and NAME not in copy
    assert "[ОСОБЫЕ ДАННЫЕ]" in copy and "Справка о сотруднике" in copy
    paragraphs = [p.text for p in Document(BytesIO(row.redacted)).paragraphs]
    assert paragraphs[1] == "[ФИО] в прошлом месяце [ОСОБЫЕ ДАННЫЕ]."
    assert not path.exists()


# ------------------------------------------------------------ оценка ---


def test_llm_detector_finds_special_and_remasks_before_trade_secret(
    real_llm: Callable[[FakeLlm], None], tmp_path: Path,
) -> None:
    (tmp_path / "a.txt").write_text(
        f"Справка.\nСотрудник в прошлом месяце [[SPECIAL_CATEGORY:{FACT}]].\n", encoding="utf-8",
    )
    sent: list[tuple[type[BaseModel], str]] = []
    real_llm(pipeline_fake(sent))
    result = evaluate_corpus(load_corpus(tmp_path), LlmDetector())
    counts = result.counts[EntityType.SPECIAL_CATEGORY]
    assert (counts.tp, counts.fn, counts.fp) == (1, 0, 0)
    assert [schema for schema, _ in sent][:1] == [SpecialCategoryAnswer]
    assert all("инфаркт" not in fragment for schema, fragment in sent[1:])


# ------------------------------- поля анкеты (разбор замера seed 24.10) ---


def test_field_label_before_masked_value_is_not_a_candidate() -> None:
    text = "Национальность: [SPECIAL_CATEGORY_1]\n[PERSON_1], инженер."
    got = [c.text for c in select_candidates(text, [], [])]
    assert got == ["[PERSON_1], инженер."]


def test_field_label_quote_before_placeholder_is_rejected(
    real_llm: Callable[[FakeLlm], None],
) -> None:
    value = "татарин"
    text = f"{NAME}, национальность: {value}, вероисповедание: ислам."
    entities = [pd_entity(text, NAME), pd_entity(text, value, EntityType.SPECIAL_CATEGORY)]
    real_llm(items("национальность", "вероисповедание: ислам"))
    found, calls = run(text, entities)
    assert [text[e.start:e.end] for e in llm_found(found)] == ["вероисповедание: ислам"]
    assert calls[0].quote == "вероисповедание: ислам"


def test_prompt_excludes_medical_grounds_without_diagnosis_and_field_labels() -> None:
    prompt = sc._system_prompt()
    assert "по медицинским показаниям" in prompt and "названия полей анкеты" in prompt
