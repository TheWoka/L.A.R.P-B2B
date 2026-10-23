"""T-031: большие документы, настоящий конвейер и API без внешних сервисов."""

import json
import logging
from collections.abc import Callable, Iterator
from io import BytesIO
from pathlib import Path
from time import monotonic
from unittest.mock import Mock

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api import documents
from app.core.config import get_settings
from app.db.models import Base, DocumentRow
from app.db.session import get_db
from app.main import app
from app.schemas import EntityType, Report, ReportText
from app.services import storage
from app.services.detection.validators import is_valid_snils
from app.services.extraction.extract import extract
from app.worker.tasks import run_pipeline
from tests.large_docs import (
    BIRTH_DATE,
    DOCX_NAME,
    DOCX_SECTIONS,
    EMAIL,
    ESTIMATE_ROWS,
    PASSPORT,
    PDF_NAME,
    PDF_PAGES,
    PHONE,
    SNILS,
    build_docx,
    build_pdf,
    build_xlsx,
)


@pytest.fixture
def large_api_db(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> Iterator[tuple[Session, Mock]]:
    """Изолированная БД; подменяется отправка в брокер, а не обработка документа."""
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(engine)
    delay = Mock()
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    monkeypatch.setenv("MAX_UPLOAD_MB", "50")
    get_settings.cache_clear()
    monkeypatch.setattr(storage, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(documents.process_document, "delay", delay)
    with Session(engine, expire_on_commit=False) as session:
        def override_db() -> Iterator[Session]:
            yield session

        app.dependency_overrides[get_db] = override_db
        try:
            yield session, delay
        finally:
            app.dependency_overrides.pop(get_db, None)
            get_settings.cache_clear()
    engine.dispose()


def assert_redacted_copy(
    client: TestClient, base: str, path: Path, names: tuple[str, ...],
) -> None:
    """Скачивает копию и повторно извлекает текст: закраски поверх ПДн недостаточно."""
    if path.suffix not in {".docx", ".pdf"}:
        return
    redacted = client.get(base + "/redacted")
    assert redacted.status_code == 200
    copy = path.with_name("redacted" + path.suffix)
    copy.write_bytes(redacted.content)
    safe_text = extract(copy, copy.name, "redacted-test").text
    private = names + ((PHONE, EMAIL) if path.suffix == ".docx" else (
        PASSPORT, SNILS, BIRTH_DATE,
    ))
    assert all(value not in safe_text for value in private)


@pytest.mark.slow
@pytest.mark.parametrize(("filename", "builder", "names", "count"), [
    ("large_contract.docx", build_docx, (DOCX_NAME,), DOCX_SECTIONS // 10),
    ("large_application.pdf", build_pdf, (PDF_NAME,), PDF_PAGES // 5),
    ("large_estimate.xlsx", build_xlsx, (), 0),
])
def test_large_document_pipeline(
    client: TestClient, large_api_db: tuple[Session, Mock], tmp_path: Path,
    filename: str, builder: Callable[[Path], None], names: tuple[str, ...], count: int,
) -> None:
    session, delay = large_api_db
    path = tmp_path / filename
    builder(path)
    response = client.post(
        "/api/documents", files={"file": (filename, path.read_bytes())},
        data={"purpose": "publication"},
    )
    assert response.status_code == 202, response.text
    task_id = response.json()["task_id"]
    delay.assert_called_once_with(task_id)
    document = session.get(DocumentRow, response.json()["document_id"])
    assert document is not None and document.stored_path is not None
    uploaded = Path(document.stored_path)
    assert uploaded.is_file()
    started = monotonic()
    run_pipeline(task_id, session)
    wall_ms = int((monotonic() - started) * 1000)
    task = client.get(f"/api/tasks/{task_id}")
    assert task.status_code == 200 and task.json()["status"] == "done", task.text
    base = f"/api/reports/{task.json()['report_id']}"
    report_response, text_response = client.get(base), client.get(base + "/text")
    assert report_response.status_code == text_response.status_code == 200
    report = Report.model_validate(report_response.json())
    text = ReportText.model_validate(text_response.json())
    assert report.overall == "red" and report.processing_ms < 120_000
    assert wall_ms < 120_000
    assert sum(e.type == EntityType.PERSON for e in report.entities) >= count
    assert report.audit.llm_provider == "mock" and not report.llm_used
    assert report.audit.llm_calls == 0
    assert not uploaded.exists() and document.stored_path is None
    for name in names:
        assert name not in text_response.text and name not in text.text
        assert all(name not in segment.text for segment in text.segments)
    if path.suffix in {".docx", ".pdf"}:
        assert report.redacted_available
    assert_redacted_copy(client, base, path, names)
    # Метрики доступны в pytest -rP; без значений ПДн.
    metrics = {
        "file": filename, "processing_ms": report.processing_ms, "worker_ms": wall_ms,
        "entities": len(report.entities), "report_bytes": len(report_response.content),
    }
    logging.getLogger(__name__).warning("T-031 metrics %s", json.dumps(metrics))


@pytest.mark.slow
def test_oversized_pdf_rejected(
    client: TestClient, large_api_db: tuple[Session, Mock],
) -> None:
    """Вход создаётся в памяти; API обязан удалить свой неполный upload."""
    _, delay = large_api_db
    payload = b"%PDF-1.7\n" + b"0" * (50 * 1024 * 1024)
    response = client.post(
        "/api/documents", files={"file": ("oversized.pdf", BytesIO(payload), "application/pdf")},
        data={"purpose": "publication"},
    )
    assert response.status_code == 400
    assert response.json() == {"detail": "Файл больше 50 МБ"}
    delay.assert_not_called()
    assert not list(storage.UPLOAD_DIR.iterdir())


@pytest.mark.slow
def test_generator_structure(tmp_path: Path) -> None:
    """Проверяет размеры корпуса и граничные разделы/страницы/строки."""
    docx, pdf, xlsx = (tmp_path / name for name in ("large.docx", "large.pdf", "large.xlsx"))
    build_docx(docx)
    document = Document(docx)
    paragraphs = document.paragraphs
    assert len(paragraphs) == DOCX_SECTIONS * 13
    assert sum(p.text.count(DOCX_NAME) for p in paragraphs) == DOCX_SECTIONS // 10
    assert sum(bool(p.paragraph_format.page_break_before) for p in paragraphs) == 69
    assert document.sections[0].header.paragraphs[0].text == "Конфиденциально"
    build_pdf(pdf)
    with pymupdf.open(pdf) as pages:
        assert len(pages) == PDF_PAGES
        assert PDF_NAME not in pages[0].get_text()
        assert PDF_NAME in pages[4].get_text() and PDF_NAME in pages[-1].get_text()
        assert sum(page.get_text().count(PDF_NAME) for page in pages) == PDF_PAGES // 5
    assert is_valid_snils(SNILS)
    build_xlsx(xlsx)
    workbook = load_workbook(xlsx, data_only=True)
    sheet = workbook.active
    assert sheet is not None and sheet.max_row == ESTIMATE_ROWS + 3
    assert sheet["B1"].value == 48300
    assert sheet["B2"].value == 0.25 and sheet["B2"].number_format == "0%"
    for row in (4, ESTIMATE_ROWS + 3):
        assert sheet.cell(row, 4).value == sheet.cell(row, 2).value * sheet.cell(row, 3).value
    workbook.close()


@pytest.mark.slow
@pytest.mark.parametrize("builder", [build_docx, build_pdf, build_xlsx])
def test_generator_missing_directory(tmp_path: Path, builder: Callable[[Path], None]) -> None:
    """Генераторы не создают отсутствующий каталог и не скрывают ошибки записи."""
    with pytest.raises((OSError, RuntimeError)):
        builder(tmp_path / "missing" / "document.bin")
