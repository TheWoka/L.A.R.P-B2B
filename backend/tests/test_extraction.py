"""Синтетические документы создаются во временной папке; бинарников в репозитории нет."""

import logging
from collections.abc import Callable
from datetime import datetime
from io import BytesIO
from pathlib import Path
from random import Random

import pymupdf
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches
from openpyxl import Workbook, load_workbook

from app.evaluation.generator.values import card, inn_person, snils
from app.schemas import EntityType, ExtractedDocument, TextSegment
from app.services.detection.detect import detect
from app.services.detection.validators import only_digits
from app.services.extraction.docx import extract_docx
from app.services.extraction.extract import extract
from app.services.extraction.pdf import extract_pdf, is_numbered_heading
from app.services.extraction.xlsx import extract_xlsx
from app.services.redaction.check import count_leftovers
from app.services.redaction.redact import redact


@pytest.fixture
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("1. Общие положения", level=1)
    document.add_paragraph("Выдуманный предмет закупки.")
    document.add_paragraph("  \t ")
    table = document.add_table(rows=2, cols=3)
    for row, values in zip(table.rows, [("Товар", "Количество", "Цена"),
                                       ("Лунный куб", "2", "10")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_heading("2. Цена контракта", level=1)
    document.add_paragraph("Условная стоимость.")
    document.save(str(path))
    return path


@pytest.fixture
def xlsx_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet.append([None, None, None])
    sheet.append(["Лунный куб", 0, 12.5])
    sheet.append([" ", None, "\t"])
    sheet.append([None, "Конец", None, "Строка"])
    sheet = workbook.create_sheet("Лист2")
    sheet.append([datetime(2025, 1, 2), False])
    sheet.append(["=1+1"])  # Без кэша формула в data_only=True даёт None.
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "sample.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "1. General terms")
        page.insert_text((72, 120), "Fictional moon cubes.\nSynthetic description.")
        page = document.new_page()
        page.insert_text((72, 72), "Continued terms.")
        page.insert_text((72, 120), "2.1. Price")
        page.insert_text((72, 170), "Synthetic price.")
        document.save(path)
    return path


@pytest.fixture
def long_pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "long.pdf"
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_htmlbox((72, 72, 500, 110), "1. Общие положения")
        page.insert_htmlbox((72, 120, 500, 160), "Выдуманный лунный куб.")
        page.insert_text((72, 200), "2. " + "a" * 117, fontsize=5)
        page.insert_text((72, 240), "Still in the first section.")
        for _ in range(50):
            document.new_page()
        document.save(path)
    return path


def test_pdf_cyrillic_long_heading_and_blank_pages(long_pdf_path: Path) -> None:
    document = extract(long_pdf_path, long_pdf_path.name, "synthetic-id")
    assert document.pages == 51
    assert "Выдуманный лунный куб." in document.text
    assert document.segments[-1].section == "1. Общие положения"
    assert all(segment.page == 1 for segment in document.segments)
    assert_offsets(document)


def assert_offsets(document: ExtractedDocument) -> None:
    assert document.text == "\n".join(segment.text for segment in document.segments)
    expected_offset = 0
    for index, segment in enumerate(document.segments):
        assert segment.id == index
        assert segment.offset == expected_offset
        assert document.text[segment.offset:segment.offset + len(segment.text)] == segment.text
        expected_offset += len(segment.text) + 1


def test_docx_order_and_sections(docx_path: Path) -> None:
    document = extract(docx_path, "sample.docx", "synthetic-doc")
    assert [segment.text for segment in document.segments] == [
        "1. Общие положения", "Выдуманный предмет закупки.",
        "Товар | Количество | Цена", "Лунный куб | 2 | 10",
        "2. Цена контракта", "Условная стоимость.",
    ]
    assert document.segments[2].section == "1. Общие положения"
    assert document.segments[-1].section == "2. Цена контракта"
    assert document.pages is None
    assert all(segment.page is None for segment in document.segments)
    assert_offsets(document)


def test_docx_headers_and_footers(tmp_path: Path) -> None:
    path = tmp_path / "margins.docx"
    source = Document()
    source.sections[0].header.paragraphs[0].text = "Коммерческая тайна. ООО «Ромашка»"
    source.sections[0].footer.paragraphs[0].text = "Срок хранения 5 лет"
    source.add_heading("1. Выдуманная закупка", level=1)
    source.add_paragraph("Поставка лунных кубов.")
    source.save(path)

    document = extract(path, path.name, "synthetic-id")
    assert document.text == (
        "Коммерческая тайна. ООО «Ромашка»\n"
        "1. Выдуманная закупка\nПоставка лунных кубов.\nСрок хранения 5 лет"
    )
    assert [segment.section for segment in document.segments] == [
        "Верхний колонтитул", "1. Выдуманная закупка",
        "1. Выдуманная закупка", "Нижний колонтитул",
    ]
    assert_offsets(document)
    segments = extract_docx(path)
    assert_offsets(document.model_copy(update={"segments": segments}))


@pytest.mark.parametrize("first_page", [False, True])
@pytest.mark.parametrize("even_pages", [False, True])
def test_docx_optional_margins_and_tables(
    tmp_path: Path, first_page: bool, even_pages: bool,
) -> None:
    path = tmp_path / "optional-margins.docx"
    source = Document()
    section = source.sections[0]
    section.different_first_page_header_footer = first_page
    source.settings.odd_and_even_pages_header_footer = even_pages
    source.add_paragraph("Выдуманное содержимое")
    expected_headers: list[str] = []
    expected_footers: list[str] = []
    for kind, expected in [("header", expected_headers), ("footer", expected_footers)]:
        for prefix, enabled in [("", True), ("first_page_", first_page),
                                ("even_page_", even_pages)]:
            margin = getattr(section, prefix + kind)
            paragraph = f"Выдуманный абзац {prefix}{kind}"
            margin.paragraphs[0].text = paragraph
            table = margin.add_table(rows=1, cols=2, width=Inches(4))
            table.cell(0, 0).text = f"Выдуманная таблица {prefix}{kind}"
            table.cell(0, 1).text = "Лунный куб"
            margin.add_paragraph(" ")
            margin.add_paragraph(f"Конец {prefix}{kind}")
            if enabled:
                expected.extend([paragraph, f"{table.cell(0, 0).text} | Лунный куб",
                                 f"Конец {prefix}{kind}"])
    source.save(path)

    document = extract(path, path.name, "synthetic-id")
    assert [segment.text for segment in document.segments] == [
        *expected_headers, "Выдуманное содержимое", *expected_footers,
    ]
    assert [segment.section for segment in document.segments] == (
        ["Верхний колонтитул"] * len(expected_headers)
        + [None] + ["Нижний колонтитул"] * len(expected_footers)
    )
    assert_offsets(document)


@pytest.mark.parametrize("linked", [False, True])
def test_docx_repeated_margins(tmp_path: Path, linked: bool) -> None:
    path = tmp_path / "repeated-margins.docx"
    source = Document()
    source.settings.odd_and_even_pages_header_footer = True
    first = source.sections[0]
    first.different_first_page_header_footer = True
    names = ["header", "first_page_header", "even_page_header",
             "footer", "first_page_footer", "even_page_footer"]
    for name in names:
        getattr(first, name).paragraphs[0].text = f"Выдуманный колонтитул {name}"
    source.add_paragraph("Первый выдуманный раздел")
    second = source.add_section()
    second.different_first_page_header_footer = True
    for name in names:
        margin = getattr(second, name)
        margin.is_linked_to_previous = linked
        if not linked:
            margin.paragraphs[0].text = f"Выдуманный колонтитул {name}"
    source.add_paragraph("Второй выдуманный раздел")
    source.save(path)

    document = extract(path, path.name, "synthetic-id")
    assert [segment.text for segment in document.segments] == [
        *(f"Выдуманный колонтитул {name}" for name in names[:3]),
        "Первый выдуманный раздел", "Второй выдуманный раздел",
        *(f"Выдуманный колонтитул {name}" for name in names[3:]),
    ]
    assert_offsets(document)


def test_xlsx_sections_and_values(xlsx_path: Path) -> None:
    document = extract(xlsx_path, "sample.xlsx", "synthetic-sheet")
    assert [segment.section for segment in document.segments] == [
        "Лист1!A2:C2", "Лист1!B4:D4", "Лист2!A1:B1",
    ]
    assert [segment.text for segment in document.segments] == [
        "Лунный куб | 0 | 12.5", "Конец | Строка", "2025-01-02 00:00:00 | False",
    ]
    assert document.pages is None
    assert all(segment.page is None for segment in document.segments)
    assert_offsets(document)


@pytest.mark.parametrize("value, number_format, expected", [
    (0.25, "0%", "25%"),
    (0.125, "0.0%", "12,5%"),
    (0, "0%", "0%"),
    (-0.125, "0.0%", "-12,5%"),
    (0.125, "0%", "13%"),
    (0.125, "0.##%", "12,5%"),
    (0.25, "0.##%", "25%"),
    (48300, "#,##0", "48300"),
    (48300, "# ##0", "48300"),
    (1234.5, "#,##0.00", "1234.5"),
    (-1234.5, "#,##0.00", "-1234.5"),
    (1234.5, "#,##0.##", "1234.5"),
    (48300, '#,##0 "₽"', "48300"),
    (48300, '0 "руб."', "48300"),
    (1234.5, '#,##0.00 "р."', "1234.5"),
    (25, '0 "%"', "25"),
    (25, r"0\%", "25"),
    (48300.0, "General", "48300"),
    (12.5, "General", "12.5"),
    (datetime(2026, 3, 12, 15, 30), "dd.mm.yyyy hh:mm", "2026-03-12 15:30:00"),
    ("Выдуманный текст", "0%", "Выдуманный текст"),
    (False, "0%", "False"),
    (None, "0%", ""),
])
def test_xlsx_cell_formats(
    tmp_path: Path, value: str | int | float | datetime | None,
    number_format: str, expected: str,
) -> None:
    path = tmp_path / "formats.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = value
    sheet["B2"].number_format = number_format
    sheet["A4"] = "Конец выдуманной таблицы"
    workbook.save(path)
    workbook.close()

    document = extract(path, path.name, "synthetic-formats")
    assert document.text == (expected + "\n" if expected else "") + "Конец выдуманной таблицы"
    assert [segment.section for segment in document.segments] == (
        (["Sheet!B2:B2"] if expected else []) + ["Sheet!A4:A4"]
    )
    assert_offsets(document)


@pytest.mark.parametrize("value, number_format, unit, expected", [
    (48300, "General", "руб.", "48300 руб."),
    (48300, "#,##0", "руб.", "48300 руб."),
    (48300, "General", "руб", "48300 руб"),
    (48300, "General", "₽", "48300 ₽"),
    (25, "General", "%", "25 %"),
    (20, "General", "шт.", "20 шт."),
    (48.3, "General", "тыс. руб.", "48.3 тыс. руб."),
    (1.2, "General", "млн руб.", "1.2 млн руб."),
    (0.25, "0%", "%", "25%"),
    (48300, '#,##0 "₽"', "руб.", "48300 руб."),
    (0, "General", "руб.", "0 руб."),
    (False, "General", "руб.", "False | руб."),
    ("48300", "General", "руб.", "48300 | руб."),
    (48300, "General", "Комментарий", "48300 | Комментарий"),
])
def test_xlsx_adjacent_units(
    tmp_path: Path, value: str | int | float, number_format: str, unit: str, expected: str,
) -> None:
    path = tmp_path / "units.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Смета"
    sheet.append(["Себестоимость единицы", value, unit, "Примечание"])
    sheet["B1"].number_format = number_format
    sheet.append([48300, None, "руб."])
    sheet.append(["руб.", 20])
    workbook.save(path)
    workbook.close()

    document = extract(path, path.name, "synthetic-units")
    assert [segment.text for segment in document.segments] == [
        f"Себестоимость единицы | {expected} | Примечание", "48300 | руб.", "руб. | 20",
    ]
    assert [segment.section for segment in document.segments] == [
        "Смета!A1:D1", "Смета!A2:C2", "Смета!A3:B3",
    ]
    assert_offsets(document)


@pytest.mark.parametrize("kind, number_format", [
    (EntityType.INN, "General"),
    (EntityType.INN, "#,##0"),
    (EntityType.SNILS, "#,##0"),
    (EntityType.CARD, "#,##0"),
    (EntityType.BIRTH_DATE, "dd.mm.yyyy"),
])
def test_xlsx_formats_preserve_pd_redaction(
    tmp_path: Path, kind: EntityType, number_format: str,
) -> None:
    rng = Random(26)
    samples = {
        EntityType.INN: ("ИНН физического лица", int(inn_person(rng))),
        EntityType.SNILS: ("СНИЛС", int(only_digits(snils(rng)))),
        EntityType.CARD: ("Номер карты", int(only_digits(card(rng)))),
        EntityType.BIRTH_DATE: ("Дата рождения", datetime(1991, 3, 12)),
    }
    label, value = samples[kind]
    path = tmp_path / "synthetic-pd.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([label, value])
    sheet["B1"].number_format = number_format
    workbook.save(path)
    workbook.close()

    document = extract(path, path.name, "synthetic-pd-formats")
    assert document.segments[0].text == f"{label} | {value}"
    entities = detect(document)
    assert any(entity.type == kind for entity in entities)
    redacted = redact(path, document, entities)
    copy = load_workbook(BytesIO(redacted.content))
    try:
        assert isinstance(copy.active["B1"].value, str)
        assert copy.active["B1"].value.startswith("[")
    finally:
        copy.close()
    assert count_leftovers(redacted, document, entities) == 0
    assert_offsets(document)


def test_xlsx_demo_estimate_trade_secrets() -> None:
    path = Path(__file__).resolve().parents[2] / "data" / "demo" / "04_estimate.xlsx"
    if not path.is_file():
        path = Path("/data/demo/04_estimate.xlsx")
    if not path.is_file():
        pytest.skip("Демо-смета отсутствует в репозитории или образе")
    document = extract(path, path.name, "synthetic-demo-estimate")
    entities = [entity for entity in detect(document) if entity.type == EntityType.TRADE_SECRET]
    for section, keyword, value in [
        ("Смета!A6:C6", "себестоимость", "48300 руб."),
        ("Смета!A7:C7", "наценка", "25%"),
    ]:
        matching = [entity for entity in entities if entity.section == section]
        assert matching, f"Не найдена коммерческая тайна: {section}"
        assert any(keyword in document.text[e.start:e.end].lower() for e in matching)
        assert any(value in document.text[e.start:e.end] for e in matching)
        assert all(e.value_masked == document.text[e.start:e.end] for e in matching)
    assert_offsets(document)


def test_pdf_pages_and_sections(pdf_path: Path) -> None:
    document = extract(pdf_path, "sample.pdf", "synthetic-pdf")
    assert document.pages == 2
    assert [segment.page for segment in document.segments] == [1, 1, 2, 2, 2]
    assert document.segments[2].section == "1. General terms"
    assert document.segments[-1].section == "2.1. Price"
    assert "Fictional moon cubes." in document.text
    assert_offsets(document)


@pytest.mark.parametrize("fixture, filename", [
    ("pdf_path", "Выдуманный_Луннокубов.PdF"),
    ("docx_path", "Выдуманный_Луннокубов.DoCx"),
    ("xlsx_path", "Выдуманный_Луннокубов.XlSx"),
])
def test_filename_dispatch_and_safe_logs(
    fixture: str, filename: str, request: pytest.FixtureRequest,
    tmp_path: Path, caplog: pytest.LogCaptureFixture,
) -> None:
    source: Path = request.getfixturevalue(fixture)
    path = tmp_path / "upload-without-extension"
    path.write_bytes(source.read_bytes())
    with caplog.at_level(logging.INFO, logger="app.services.extraction.extract"):
        document = extract(path, filename, "synthetic-id")
    assert document.text
    assert document.filename == filename
    assert document.document_id == "synthetic-id"
    assert filename not in caplog.text
    assert Path(filename).stem not in caplog.text
    assert "synthetic-id" in caplog.text
    assert f"сегментов={len(document.segments)}" in caplog.text
    assert f"страниц={document.pages}" in caplog.text
    assert all(segment.text not in caplog.text for segment in document.segments)
    assert_offsets(document)


def test_unsupported_extension(tmp_path: Path) -> None:
    with pytest.raises(ValueError) as error:
        extract(tmp_path / "missing", "sample.txt", "synthetic-id")
    assert str(error.value) == (
        "Формат файла не поддерживается. Загрузите .pdf, .docx или .xlsx"
    )


@pytest.fixture(params=["pdf", "docx", "xlsx"])
def corrupt_path(tmp_path: Path, request: pytest.FixtureRequest) -> Path:
    extension: str = request.param
    path = tmp_path / f"broken.{extension}"
    path.write_bytes(b"\x00\xff\x13\x89not-a-document")
    return path


def test_corrupt_file(corrupt_path: Path) -> None:
    with pytest.raises(ValueError, match="Файл повреждён") as error:
        extract(corrupt_path, corrupt_path.name, "synthetic-id")
    assert error.value.__cause__ is None


@pytest.fixture(params=["docx", "xlsx"])
def empty_path(tmp_path: Path, request: pytest.FixtureRequest) -> Path:
    extension: str = request.param
    path = tmp_path / f"empty.{extension}"
    if extension == "docx":
        document = Document()
        document.add_paragraph(" \t ")
        document.add_table(rows=1, cols=2)
        document.save(str(path))
    else:
        workbook = Workbook()
        workbook.active.append([None, " \t "])
        workbook.save(path)
        workbook.close()
    return path


def test_empty_documents(empty_path: Path) -> None:
    with pytest.raises(ValueError, match="^В документе нет текста для проверки\\.$"):
        extract(empty_path, empty_path.name, "synthetic-id")


@pytest.mark.parametrize("page_count", [1, 2])
def test_pdf_without_text_reports_scan(tmp_path: Path, page_count: int) -> None:
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as source:
        for _ in range(page_count):
            page = source.new_page()
            page.insert_text((72, 72), "   ")
        source.save(path)

    with pytest.raises(ValueError) as error:
        extract(path, path.name, "synthetic-id")
    assert str(error.value) == (
        "Документ похож на скан: текст не распознаётся. "
        "Загрузите документ с текстовым слоем."
    )


@pytest.mark.parametrize("empty_pages", [(1,), (2,), (1, 3)])
def test_pdf_partially_empty_pages_warning(
    tmp_path: Path, caplog: pytest.LogCaptureFixture, empty_pages: tuple[int, ...],
) -> None:
    path = tmp_path / "partially-empty.pdf"
    page_count = len(empty_pages) + 1
    with pymupdf.open() as source:
        for number in range(1, page_count + 1):
            page = source.new_page()
            if number not in empty_pages:
                page.insert_text((72, 72), "Fictional moon cubes.")
        source.save(path)

    with caplog.at_level(logging.WARNING, logger="app.services.extraction.pdf"):
        document = extract(path, path.name, "synthetic-id")
    assert document.pages == page_count
    assert document.text == "Fictional moon cubes."
    assert [segment.page for segment in document.segments] == [
        number for number in range(1, page_count + 1) if number not in empty_pages
    ]
    warnings = [record for record in caplog.records if record.levelno == logging.WARNING]
    assert len(warnings) == 1
    assert warnings[0].getMessage() == f"PDF: страницы без текста: {list(empty_pages)}"
    assert document.text not in caplog.text
    assert_offsets(document)


@pytest.fixture
def merged_path(tmp_path: Path) -> Path:
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Общий заголовок"
    table.cell(0, 2).text = "Правый столбец"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Общая категория"
    table.cell(1, 1).text = "Одинаково"
    table.cell(1, 2).text = "Одинаково"
    table.cell(2, 2).text = "Конец"
    document.save(str(path))
    return path


def test_docx_merged_cells(merged_path: Path) -> None:
    extracted = extract(merged_path, merged_path.name, "synthetic-id")
    assert [segment.text for segment in extracted.segments] == [
        "Общий заголовок | Правый столбец",
        "Общая категория | Одинаково | Одинаково", "Конец",
    ]
    assert_offsets(extracted)


@pytest.fixture(params=[
    ("Без номера", "Heading 2", "Без номера"),
    ("Русский стиль", "Заголовок 2", "Русский стиль"),
    ("2026 год — отчёт", "Heading 2", "2026 год — отчёт"),
    ("3.1. Цена контракта составляет 100 000 руб.", "Заголовок 2",
     "3.1. Цена контракта составляет 100 000 руб."),
    ("а" * 81 + ":", "Heading 2", "а" * 81 + ":"),
    ("3.2. Условия", "Normal", "3.2. Условия"),
    ("3. " + "а" * 117, "Normal", None),
    ("Обычный абзац", "Normal", None),
])
def heading_case(tmp_path: Path, request: pytest.FixtureRequest) -> tuple[Path, str | None]:
    heading, style, expected = request.param
    path = tmp_path / "headings.docx"
    document = Document()
    if style not in document.styles:
        document.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph(heading, style)
    document.add_paragraph("Текст раздела")
    document.save(str(path))
    return path, expected


def test_docx_heading_detection(heading_case: tuple[Path, str | None]) -> None:
    path, expected = heading_case
    assert extract_docx(path)[-1].section == expected


@pytest.mark.parametrize("line, expected", [
    ("1. Общие положения", True),
    ("3.2. Обязанности сторон", True),
    ("2026 год — отчёт", False),
    ("3.1. Цена контракта составляет 100 000 руб.", False),
    ("10 000 рублей", False),
    ("1.1.1. Термины", True),
    ("99.99.99 Термины", True),
    ("1.1.1.1. Термины", False),
    ("100. Термины", False),
    ("1.100. Термины", False),
    ("1.1.100. Термины", False),
    ("1. термины", False),
    ("1.Общие положения", False),
    ("1. Термины;", False),
    ("1. Термины:", False),
    ("1. Термины,", False),
    ("1. " + "А" * 77, True),
    ("1. " + "А" * 78, False),
    ("1. Термины" + " слова" * 8, True),
    ("1. Термины" + " слова" * 9, False),
    ("", False),
])
@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_numbered_heading_rules(
    tmp_path: Path, line: str, expected: bool, extension: str,
) -> None:
    assert is_numbered_heading(line) is expected
    path = tmp_path / f"numbered-heading.{extension}"
    if extension == "docx":
        source = Document()
        source.add_paragraph(line)
        source.add_paragraph("Выдуманный текст раздела")
        source.save(path)
    else:
        with pymupdf.open() as source:
            page = source.new_page(width=1400)
            if line:
                page.insert_htmlbox((40, 40, 1360, 100), line)
            page.insert_htmlbox((40, 150, 1360, 200), "Выдуманный текст раздела")
            source.save(path)

    document = extract(path, path.name, "synthetic-id")
    assert document.segments[-1].section == (line if expected else None)
    assert_offsets(document)


@pytest.mark.parametrize("fixture, parser", [
    ("docx_path", extract_docx), ("xlsx_path", extract_xlsx), ("pdf_path", extract_pdf),
])
def test_parser_offsets(
    fixture: str, parser: Callable[[Path], list[TextSegment]], request: pytest.FixtureRequest,
) -> None:
    path: Path = request.getfixturevalue(fixture)
    segments = parser(path)
    assert segments
    assert_offsets(ExtractedDocument(
        document_id="synthetic-id", filename=path.name,
        text="\n".join(segment.text for segment in segments), segments=segments,
    ))


def test_extract_reindexes_segments(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_parser(path: Path) -> list[TextSegment]:
        return [TextSegment(id=8, offset=99, text="Первый\nблок"),
                TextSegment(id=5, offset=20, text=" \t "),
                TextSegment(id=3, offset=0, text="Повтор"),
                TextSegment(id=4, offset=0, text="Повтор")]

    monkeypatch.setattr("app.services.extraction.extract.extract_docx", fake_parser)
    document = extract(tmp_path / "unused", "synthetic.docx", "synthetic-id")
    assert len(document.segments) == 3
    assert_offsets(document)
