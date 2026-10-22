"""Обезличенная копия (T-026): форматы, самопроверка, конвейер, API и срок хранения.

Только выдуманные данные; LLM_PROVIDER=mock. Документы создаются в tmp_path.
"""

import logging
from collections.abc import Iterator
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from random import Random
from unittest.mock import Mock
from urllib.parse import quote

import pymupdf
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api import documents, reports
from app.core.config import get_settings
from app.db.models import Base, DocumentRow, ReportRow, TaskRow
from app.db.session import get_db
from app.evaluation.generator.values import inn_person
from app.main import app
from app.schemas import Entity, EntityCategory, EntityType, ExtractedDocument, Report, TextSegment
from app.services import retention, storage
from app.services.detection.detect import detect
from app.services.extraction.extract import extract
from app.services.redaction.check import copy_texts, count_leftovers
from app.services.redaction.redact import RedactedFile, redact
from app.services.redaction.values import (
    LABELS,
    PdValue,
    normalize_spaces,
    pd_values,
    replace_values,
)
from app.worker import tasks

# Синтетические реквизиты; СНИЛС и ИНН — с корректными контрольными суммами.
NAME = "Иванов Иван Иванович"
HEADER_NAME = "Петров Пётр Петрович"
TABLE_NAME = "Смирнова Анна Сергеевна"
SPLIT_NAME = "Сидоров Сергей Сергеевич"
PHONE = "+7 (912) 345-67-89"
SNILS = "112-233-445 95"
EMAIL = "ivan.petrov@example.ru"
INN = inn_person(Random(26))
COST = "Себестоимость единицы 48 300 руб."
TEXT = (
    f"Документ для публикации. Контакт: {NAME}, телефон {PHONE}.\n"
    f"СНИЛС: {SNILS}.\n{COST}\nПовторно: {NAME}."
)


def detected(path: Path, filename: str | None = None) -> tuple[ExtractedDocument, list[Entity]]:
    """Извлечение и настоящая детекция, как в конвейере."""
    if path.suffix == ".txt":
        text = path.read_text(encoding="utf-8")
        doc = ExtractedDocument(document_id="synthetic", filename="Текстовый запрос", text=text,
                                segments=[TextSegment(id=0, text=text, offset=0)])
    else:
        doc = extract(path, filename or path.name, "synthetic")
    return doc, detect(doc)


def assert_clean_copy(
    redacted: RedactedFile, values: tuple[str, ...], labels: tuple[str, ...],
) -> str:
    """В тексте копии нет исходных ПДн, есть метки, коммерческая тайна на месте."""
    text = normalize_spaces(copy_texts(redacted, "synthetic")[0])
    for value in values:
        assert value not in text
    for label in labels:
        assert label in text
    assert COST in text
    return text


def entity_types(entities: list[Entity]) -> set[EntityType]:
    return {entity.type for entity in entities}


# ---------- Форматы ----------
def test_txt_redaction(tmp_path: Path) -> None:
    path = tmp_path / "query.txt"
    path.write_text(TEXT, encoding="utf-8")
    doc, entities = detected(path)
    assert entity_types(entities) >= {EntityType.PERSON, EntityType.PHONE, EntityType.SNILS}
    redacted = redact(path, doc, entities)
    assert redacted.ext == "txt"
    text = assert_clean_copy(redacted, (NAME, PHONE, SNILS), ("[ФИО]", "[ТЕЛЕФОН]", "[СНИЛС]"))
    assert text.count("[ФИО]") == 2
    assert count_leftovers(redacted, doc, entities) == 0


def add_hyperlink(paragraph_part: object, paragraph: object, address: str) -> None:
    """Ссылка mailto: — значение и в тексте run внутри w:hyperlink, и в связях части."""
    r_id = paragraph_part.relate_to(f"mailto:{address}", RELATIONSHIP_TYPE.HYPERLINK,
                                    is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = address
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def make_docx(path: Path) -> None:
    document = Document()
    document.core_properties.author = HEADER_NAME
    document.core_properties.title = f"Заявка {NAME}"
    document.sections[0].header.paragraphs[0].text = f"Исполнитель: {HEADER_NAME}"
    document.sections[0].footer.paragraphs[0].text = f"Телефон для связи: {PHONE}"
    document.add_paragraph("Техническое задание для публикации")
    document.add_paragraph(f"Контакт: {NAME}, СНИЛС {SNILS}.")
    split = document.add_paragraph("Представитель: ")
    split.add_run(SPLIT_NAME[:10])
    split.add_run(SPLIT_NAME[10:]).bold = True
    split.add_run(", действует по доверенности.")
    mail = document.add_paragraph("Почта: ")
    add_hyperlink(document.part, mail, EMAIL)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Контактное лицо"
    table.rows[0].cells[1].text = TABLE_NAME
    document.add_paragraph(COST)
    document.save(path)


def test_docx_redaction(tmp_path: Path) -> None:
    path = tmp_path / "tz.docx"
    make_docx(path)
    doc, entities = detected(path)
    values = (NAME, HEADER_NAME, TABLE_NAME, SPLIT_NAME, PHONE, SNILS, EMAIL)
    found = {doc.text[entity.start:entity.end] for entity in entities}
    assert set(values) <= found, "детектор должен найти все синтетические ПДн"
    redacted = redact(path, doc, entities)
    assert redacted.ext == "docx"
    assert_clean_copy(redacted, values, ("[ФИО]", "[ТЕЛЕФОН]", "[СНИЛС]", "[EMAIL]"))
    assert count_leftovers(redacted, doc, entities) == 0
    copy = Document(BytesIO(redacted.content))
    assert copy.sections[0].header.paragraphs[0].text == "Исполнитель: [ФИО]"
    assert copy.tables[0].rows[0].cells[1].text == "[ФИО]"
    split = copy.paragraphs[2]
    assert split.text == "Представитель: [ФИО], действует по доверенности."
    assert [run.text for run in split.runs] == [
        "Представитель: ", "[ФИО]", "", ", действует по доверенности.",
    ]
    assert split.runs[2].bold, "форматирование затронутого run сохраняется"
    assert copy.paragraphs[-1].text == COST, "абзац без ПДн не меняется"
    assert copy.core_properties.author == "" and NAME not in copy.core_properties.title
    targets = [rel.target_ref for rel in copy.part.rels.values() if rel.is_external]
    assert targets == ["mailto:[EMAIL]"]


def test_xlsx_redaction(tmp_path: Path) -> None:
    path = tmp_path / "estimate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Смета"
    rows = [("ФИО исполнителя", NAME), ("Телефон", PHONE), ("ИНН физического лица", int(INN)),
            (COST, None), ("Количество", 20), ("Сумма", "=B5*48300")]
    for row in rows:
        sheet.append(row)
    workbook.properties.creator = NAME
    workbook.save(path)
    doc, entities = detected(path)
    assert entity_types(entities) >= {EntityType.PERSON, EntityType.PHONE, EntityType.INN}
    redacted = redact(path, doc, entities)
    assert redacted.ext == "xlsx"
    assert_clean_copy(redacted, (NAME, PHONE, INN), ("[ФИО]", "[ТЕЛЕФОН]", "[ИНН]"))
    assert count_leftovers(redacted, doc, entities) == 0
    copy = load_workbook(BytesIO(redacted.content))["Смета"]
    assert copy["B3"].value == "[ИНН]", "число целиком заменено строкой-меткой"
    assert copy["B5"].value == 20 and copy["B6"].value == "=B5*48300"
    assert NAME not in str(copy.parent.properties.creator)  # openpyxl пишет «openpyxl»


def make_pdf(path: Path, lines: list[str]) -> None:
    document = pymupdf.open()
    page = document.new_page()
    font = pymupdf.Font("helv")
    page.insert_font(fontname="F_test", fontbuffer=font.buffer)
    for number, line in enumerate(lines):
        page.insert_text((72, 100 + 24 * number), line, fontname="F_test", fontsize=11)
    document.set_metadata({"author": NAME, "title": "Заявка"})
    document.save(path)


def test_pdf_redaction_removes_text(tmp_path: Path) -> None:
    path = tmp_path / "application.pdf"
    make_pdf(path, ["Заявка участника", f"Представитель: {NAME}.", f"Телефон: {PHONE}.",
                    f"СНИЛС: {SNILS}.", COST])
    doc, entities = detected(path)
    assert entity_types(entities) >= {EntityType.PERSON, EntityType.PHONE, EntityType.SNILS}
    redacted = redact(path, doc, entities)
    assert redacted.ext == "pdf"
    assert_clean_copy(redacted, (NAME, PHONE, SNILS), ("[ФИО]", "[ТЕЛЕФОН]", "[СНИЛС]"))
    assert count_leftovers(redacted, doc, entities) == 0
    with pymupdf.open(stream=redacted.content, filetype="pdf") as copy:
        page = copy[0]
        raw = page.get_text()
        for value in (NAME, PHONE, SNILS, "Иванов", "345-67-89"):
            assert value not in raw
            assert page.search_for(value) == [], "текст удалён, а не закрашен"
        assert page.search_for("48 300")
        assert NAME not in str(copy.metadata)


def test_pdf_value_on_other_page_is_removed(tmp_path: Path) -> None:
    """Сущность отмечена на странице 1, повтор на странице 2 тоже удаляется."""
    path = tmp_path / "two-pages.pdf"
    make_pdf(path, [f"Представитель: {NAME}."])
    with pymupdf.open(path) as document:
        page = document.new_page()
        page.insert_font(fontname="F_test", fontbuffer=pymupdf.Font("helv").buffer)
        page.insert_text((72, 100), f"Подпись: {NAME}", fontname="F_test", fontsize=11)
        data = document.tobytes()
    path.write_bytes(data)
    doc = extract(path, path.name, "synthetic")
    start = doc.text.index(NAME)
    entity = Entity(type=EntityType.PERSON, category=EntityCategory.pd, start=start,
                    end=start + len(NAME), page=1, value_masked="И", confidence=1,
                    detector="ner")
    redacted = redact(path, doc, [entity])
    assert NAME not in copy_texts(redacted, "synthetic")[0]


# ---------- Значения и граничные случаи ----------
def test_short_value_only_as_whole_word() -> None:
    values = [PdValue(text="Ли", label="[ФИО]")]
    assert replace_values("Лист согласования. Подписал Ли.", values) == (
        "Лист согласования. Подписал [ФИО]."
    )


def test_pd_values_split_by_lines_and_cells() -> None:
    text = "Представитель: Иванов Иван\nИванович | Сидоров Сергей"
    doc = ExtractedDocument(document_id="x", filename="x.pdf", text=text,
                            segments=[TextSegment(id=0, text=text, offset=0)])
    entity = Entity(type=EntityType.PERSON, category=EntityCategory.pd, start=15,
                    end=len(text), value_masked="И", confidence=1, detector="ner")
    trade = Entity(type=EntityType.TRADE_SECRET, category=EntityCategory.trade_secret,
                   start=0, end=14, value_masked="x", confidence=1, detector="dict")
    assert [value.text for value in pd_values(doc, [entity, trade])] == [
        "Сидоров Сергей", "Иванов Иван", "Иванович",
    ]


def test_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "file.odt"
    path.write_bytes(b"synthetic")
    doc = ExtractedDocument(document_id="x", filename="file.odt", text="", segments=[])
    with pytest.raises(ValueError, match="не поддерживается"):
        redact(path, doc, [])


def test_self_check_finds_value_left_in_docx(tmp_path: Path) -> None:
    """Исходник вместо копии: самопроверка находит все значения ПДн."""
    path = tmp_path / "tz.docx"
    make_docx(path)
    doc, entities = detected(path)
    original = RedactedFile(content=path.read_bytes(), ext="docx")
    assert count_leftovers(original, doc, entities) >= 6


# ---------- Конвейер (worker) ----------
@pytest.fixture
def worker_db() -> Iterator[Session]:
    engine = create_engine("sqlite://", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as session:
        yield session
    engine.dispose()


def run_text_task(session: Session, tmp_path: Path) -> tuple[TaskRow, Path]:
    path = tmp_path / "query.txt"
    path.write_text(TEXT, encoding="utf-8")
    session.add(DocumentRow(id="doc", filename=None, purpose="publication", stored_path=str(path)))
    session.flush()
    task = TaskRow(id="task", document_id="doc", status="queued")
    session.add(task)
    session.commit()
    tasks.run_pipeline(task.id, session)
    return task, path


def test_worker_stores_copy(worker_db: Session, tmp_path: Path) -> None:
    task, path = run_text_task(worker_db, tmp_path)
    assert task.status == "done"
    row = worker_db.get(ReportRow, task.report_id)
    assert Report.model_validate(row.data).redacted_available
    assert row.redacted_ext == "txt" and NAME not in row.redacted.decode("utf-8")
    assert not path.exists()


def test_self_check_failure_hides_copy(
    worker_db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    leaky = Mock(return_value=RedactedFile(content=f"Контакт: {NAME}".encode(), ext="txt"))
    monkeypatch.setattr(tasks, "redact", leaky)
    with caplog.at_level(logging.INFO):
        task, path = run_text_task(worker_db, tmp_path)
    leaky.assert_called_once()
    assert task.status == "done"
    row = worker_db.get(ReportRow, task.report_id)
    assert not Report.model_validate(row.data).redacted_available
    assert row.redacted is None and row.redacted_ext is None
    assert "redaction self-check failed: ext=txt remaining_values=1" in caplog.text
    assert NAME not in caplog.text and "Иванов" not in caplog.text
    assert not path.exists()


def test_redaction_error_keeps_report(
    worker_db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    monkeypatch.setattr(tasks, "redact", Mock(side_effect=RuntimeError(NAME)))
    with caplog.at_level(logging.INFO):
        task, path = run_text_task(worker_db, tmp_path)
    assert (task.status, task.progress, task.error) == ("done", 100, None)
    row = worker_db.get(ReportRow, task.report_id)
    assert not Report.model_validate(row.data).redacted_available and row.redacted is None
    assert "redaction failed exception_type=RuntimeError" in caplog.text
    assert NAME not in caplog.text
    assert not path.exists()


# ---------- API, сквозной сценарий ----------
@pytest.fixture
def api_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Session]:
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(engine)
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    get_settings.cache_clear()
    monkeypatch.setattr(storage, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(documents.process_document, "delay", Mock())
    with Session(engine, expire_on_commit=False) as session:
        app.dependency_overrides[get_db] = lambda: session
        try:
            yield session
        finally:
            app.dependency_overrides.pop(get_db, None)
            get_settings.cache_clear()
    engine.dispose()


def upload_and_process(client: TestClient, session: Session, path: Path) -> str:
    response = client.post("/api/documents", files={"file": (path.name, path.read_bytes())},
                           data={"purpose": "publication"})
    assert response.status_code == 202
    task_id = response.json()["task_id"]
    tasks.run_pipeline(task_id, session)
    task = client.get(f"/api/tasks/{task_id}").json()
    assert task["status"] == "done", task
    return task["report_id"]


def expected_disposition(name: str, ext: str) -> str:
    return f"attachment; filename=\"redacted.{ext}\"; filename*=UTF-8''{quote(name, safe='')}"


def test_e2e_docx_with_pd_in_filename(
    client: TestClient, api_db: Session, tmp_path: Path,
) -> None:
    path = tmp_path / "Иванов_Иван_Иванович_заявка.docx"
    make_docx(path)
    report_id = upload_and_process(client, api_db, path)
    report = Report.model_validate(client.get(f"/api/reports/{report_id}").json())
    assert report.redacted_available and report.overall == "red"
    response = client.get(f"/api/reports/{report_id}/redacted")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.headers["content-disposition"] == expected_disposition(
        "документ_обезличено.docx", "docx",
    )
    copy = Document(BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in copy.paragraphs)
    for value in (NAME, SPLIT_NAME, SNILS, EMAIL, "Иванов"):
        assert value not in text and value.encode() not in response.content
    assert COST in text and "[ФИО]" in text
    assert not list(storage.UPLOAD_DIR.iterdir()), "исходник удалён"
    assert api_db.get(DocumentRow, report.document_id).stored_path is None


def test_e2e_text_query_and_plain_filename(
    client: TestClient, api_db: Session, tmp_path: Path,
) -> None:
    response = client.post("/api/queries", json={"text": TEXT, "purpose": "publication"})
    task_id = response.json()["task_id"]
    tasks.run_pipeline(task_id, api_db)
    report_id = client.get(f"/api/tasks/{task_id}").json()["report_id"]
    response = client.get(f"/api/reports/{report_id}/redacted")
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/plain; charset=utf-8"
    assert response.headers["content-disposition"] == expected_disposition(
        "запрос_обезличено.txt", "txt",
    )
    assert NAME not in response.text and COST in response.text


@pytest.mark.parametrize(("filename", "stem"), [
    ("ТЗ_поставка.docx", "ТЗ_поставка"),
    ("Иванов_заявка.docx", "документ"),
    ("смета +7 912 345-67-89.xlsx", "документ"),
    (".pdf", "документ"),
    (None, "запрос"),
])
def test_download_stem(filename: str | None, stem: str) -> None:
    assert reports.download_stem(filename) == stem


def test_report_without_copy_returns_404(client: TestClient, api_db: Session) -> None:
    api_db.add(DocumentRow(id="doc", purpose="internal"))
    api_db.flush()
    api_db.add(ReportRow(id="old", document_id="doc", data={"filename": None}))
    api_db.commit()
    response = client.get("/api/reports/old/redacted")
    assert response.status_code == 404
    assert response.json() == {"detail": "Обезличенная копия недоступна"}


# ---------- Срок хранения ----------
def test_retention_removes_copy_with_report(
    api_db: Session, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("REPORT_TTL_DAYS", "7")
    get_settings.cache_clear()
    now = datetime(2026, 10, 22, 12, tzinfo=timezone.utc)
    for identifier, days in (("expired", 10), ("current", 1)):
        api_db.add(DocumentRow(id=identifier, purpose="internal"))
        api_db.flush()
        api_db.add(ReportRow(id=identifier, document_id=identifier, data={},
                             redacted=b"synthetic copy", redacted_ext="txt",
                             created_at=now - timedelta(days=days)))
    api_db.commit()
    assert retention.purge_expired(api_db, now) == 1
    rows = api_db.execute(select(ReportRow.id, ReportRow.redacted)).all()
    assert rows == [("current", b"synthetic copy")]


# ---------- Демо-набор ----------
DEMO_DIR = Path(__file__).resolve().parents[2] / "data" / "demo"
DEMO_FILES = ("01_tz_publication.docx", "02_contract_clean.docx", "03_application_pd.pdf",
              "04_estimate.xlsx", "05_anketa_special.docx", "06_medical_tz.docx")


@pytest.mark.parametrize("filename", DEMO_FILES)
def test_demo_documents_pass_self_check(filename: str) -> None:
    path = DEMO_DIR / filename
    if not path.is_file():
        pytest.skip(f"{path} недоступен (в образе api нет data/demo)")
    doc, entities = detected(path)
    redacted = redact(path, doc, entities)
    assert redacted.ext == path.suffix.lstrip(".")
    assert count_leftovers(redacted, doc, entities) == 0
    pd = [entity for entity in entities
          if entity.category in (EntityCategory.pd, EntityCategory.pd_special)]
    text = copy_texts(redacted, "demo")[0]
    for entity in pd:
        assert doc.text[entity.start:entity.end] not in text
    assert any(label in text for label in LABELS.values()) == bool(pd)
