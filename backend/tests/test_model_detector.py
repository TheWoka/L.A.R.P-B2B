"""Классификатор grif-cls в конвейере (T-036b). Только выдуманные данные.

Тесты с настоящей моделью пропускаются, если data/models/grif-cls недоступна
(в образе api модель приходит томом ./data/models:/data/models:ro из docker-compose.yml).
Логика пересечений проверяется на фиктивной модели — без весов.
"""

import logging
from collections.abc import Iterator
from io import BytesIO
from pathlib import Path
from time import monotonic

import pytest
import yaml
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.models import Base, DocumentRow, ReportRow, TaskRow
from app.evaluation.report import LlmDetector, detect_with_model
from app.schemas import (
    DocPurpose,
    Entity,
    EntityCategory,
    EntityType,
    ExtractedDocument,
    Report,
    ReportText,
    TextSegment,
)
from app.services.detection import model_detector as md
from app.services.detection.detect import detect
from app.services.detection.fragments import split_fragments
from app.services.extraction.extract import extract
from app.services.llm.special_category import has_person_context
from app.services.masking.masker import mask
from app.services.pipeline import run_audit
from app.worker.tasks import run_pipeline
from tests.large_docs import build_docx

# Вымышленные люди; совпадения с реальными случайны.
MANAGER = "Кукушкин Илья Петрович"
EMPLOYEE = "Зябликов Аркадий Венедиктович"
TS_TEXT = f"Менеджер {MANAGER} сообщил: берём у завода по 900 руб., клиенту отдаём по 1 200 руб."
TS_PHRASE = "берём у завода по 900 руб., клиенту отдаём по 1 200 руб."
SPECIAL_TEXT = f"Сотрудник {EMPLOYEE} в прошлом месяце перенёс операцию на сердце."
SPECIAL_PHRASE = "в прошлом месяце перенёс операцию на сердце"
NO_PERSON_TEXT = "Поставка препаратов для лечения онкологических заболеваний для больницы № 3."
# Фраза из data/demo/06_medical_tz.docx: модель даёт SPECIAL_CATEGORY с p ≈ 0,93.
MEDICAL_TEXT = ("Номенклатура включает препараты для лечения онкологических заболеваний "
                "и средства для пациентов с диабетом.")
DEMO_DIR = Path(__file__).resolve().parents[2] / "data" / "demo"

requires_model = pytest.mark.skipif(md.find_model_dir() is None,
                                    reason="data/models/grif-cls недоступна")


def make_doc(text: str) -> ExtractedDocument:
    return ExtractedDocument(document_id="synthetic-model", filename="synthetic.txt", text=text,
                             segments=[TextSegment(id=0, text=text, offset=0)])


def model_found(text: str) -> list[Entity]:
    return [e for e in detect_with_model(make_doc(text)) if e.detector == "model"]


@pytest.fixture
def no_model(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Iterator[None]:
    monkeypatch.setattr(md, "MODEL_DIRS", (tmp_path / "missing",))
    md.load_model.cache_clear()
    yield
    md.load_model.cache_clear()


# ---------------------------------------------------- настоящая модель ---


@requires_model
def test_trade_secret_without_keywords() -> None:
    found = model_found(TS_TEXT)
    assert [e.type for e in found] == [EntityType.TRADE_SECRET]
    secret = found[0]
    assert TS_PHRASE in TS_TEXT[secret.start:secret.end]
    assert MANAGER not in TS_TEXT[secret.start:secret.end]       # ФИО вырезано из фрагмента
    assert secret.value_masked == TS_TEXT[secret.start:secret.end]  # КТ — без маскирования
    assert secret.category == EntityCategory.trade_secret
    assert 0.35 <= secret.confidence <= 1


@requires_model
def test_special_category_masked_and_cut_from_person() -> None:
    doc = make_doc(SPECIAL_TEXT)
    result = run_audit(doc, DocPurpose.internal, SPECIAL_TEXT.encode())
    special = [e for e in result.report.entities if e.type == EntityType.SPECIAL_CATEGORY]
    person = [e for e in result.report.entities if e.type == EntityType.PERSON]
    assert len(special) == 1 and special[0].detector == "model" and len(person) == 1
    s, p = special[0], person[0]
    assert s.category == EntityCategory.pd_special
    assert not (s.start < p.end and p.start < s.end)            # не пересекается с PERSON
    assert SPECIAL_TEXT[s.start:s.end].startswith(SPECIAL_PHRASE)
    shown = result.report_text.text[s.start:s.end]
    assert "операцию" not in shown and "сердце" not in shown    # в отчёте замаскировано
    assert s.value_masked == shown and len(shown) == s.end - s.start


@requires_model
def test_medical_procurement_without_people_is_not_special() -> None:
    assert model_found(NO_PERSON_TEXT) == []
    assert model_found(MEDICAL_TEXT) == []                       # p ≈ 0,93, но человека нет


@requires_model
def test_person_filter_decides_for_medical_phrase(monkeypatch: pytest.MonkeyPatch) -> None:
    """Без привязки к человеку та же фраза стала бы особой категорией — фильтр и решает."""
    with_person = f"{MEDICAL_TEXT} Ответственный — {EMPLOYEE}."
    found = detect_with_model(make_doc(with_person))         # словарь или модель — не дубль
    assert [e.type for e in found].count(EntityType.SPECIAL_CATEGORY) == 1
    monkeypatch.setattr(md, "has_person_context",
                        lambda texts, **_: [True] * len(texts))
    assert [e.type for e in model_found(MEDICAL_TEXT)] == [EntityType.SPECIAL_CATEGORY]


@requires_model
def test_no_duplicate_of_dictionary_trade_secret() -> None:
    text = "Наценка на кабель составит 14%."
    entities = detect_with_model(make_doc(text))
    assert [(e.type, e.detector) for e in entities] == [(EntityType.TRADE_SECRET, "dict")]


@requires_model
def test_audit_detectors_contain_model_version() -> None:
    result = run_audit(make_doc(TS_TEXT), DocPurpose.internal, TS_TEXT.encode())
    version = md.model_version()
    assert version is not None and version.startswith("grif-cls-") and len(version) == 17
    assert f"model:{version}" in result.report.audit.detectors
    assert any(e.detector == "model" for e in result.report.entities)


@requires_model
def test_llm_detector_order_includes_model() -> None:
    """--llm в mock: тот же результат, что detect_with_model (LLM-шаги ничего не делают)."""
    doc = make_doc(TS_TEXT)
    assert LlmDetector()(doc) == detect_with_model(doc)
    assert LlmDetector(use_model=False)(doc) == detect(doc)


@pytest.fixture
def session() -> Iterator[Session]:
    engine = create_engine("sqlite://", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as db:
        yield db
    engine.dispose()


@requires_model
def test_special_category_cut_from_redacted_docx(session: Session, tmp_path: Path) -> None:
    path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("Служебная записка")
    document.add_paragraph(SPECIAL_TEXT)
    document.save(path)
    session.add(DocumentRow(id="synthetic-document", filename="synthetic.docx",
                            purpose="internal", stored_path=str(path)))
    session.flush()
    task = TaskRow(id="synthetic-task", document_id="synthetic-document", status="queued")
    session.add(task)
    session.commit()
    run_pipeline(task.id, session)
    assert task.status == "done"
    row = session.get(ReportRow, task.report_id)
    report, text = Report.model_validate(row.data), ReportText.model_validate(row.text)
    assert any(e.type == EntityType.SPECIAL_CATEGORY and e.detector == "model"
               for e in report.entities)
    assert "операцию на сердце" not in text.text
    assert report.redacted_available and row.redacted_ext == "docx"
    copy = "\n".join(p.text for p in Document(BytesIO(row.redacted)).paragraphs)
    assert "операцию на сердце" not in copy and "Зябликов" not in copy
    assert "Служебная записка" in copy


@requires_model
@pytest.mark.parametrize("item", yaml.safe_load((DEMO_DIR / "answers.yaml").read_text(
    encoding="utf-8")) if (DEMO_DIR / "answers.yaml").is_file() else [None])
def test_demo_regression(item: dict | None) -> None:
    """Все 6 демо в mock: ровно overall и набор нарушений из answers.yaml."""
    if item is None:
        pytest.skip(f"{DEMO_DIR} недоступен (в образе api нет data/demo)")
    path = DEMO_DIR / item["file"]
    doc = extract(path, path.name, "synthetic-demo")
    result = run_audit(doc, DocPurpose(item["purpose"]), path.read_bytes())
    assert result.report.overall == item["overall"]
    assert sorted(v.rule_id for v in result.report.violations) == sorted(item["violations"])


@requires_model
@pytest.mark.slow
def test_model_step_speed_on_70_pages(tmp_path: Path) -> None:
    path = tmp_path / "large_contract.docx"
    build_docx(path)
    doc = extract(path, path.name, "synthetic-large")
    entities = detect(doc)
    masked = mask(doc.text, entities)
    md.load_model()                                              # загрузка — не часть шага
    started = monotonic()
    md.apply_model(doc, entities, masked)
    assert monotonic() - started < 10


# ------------------------------------------------------- модели нет ---


def test_missing_model_skips_step_with_one_warning(
    no_model: None, caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.WARNING, logger=md.__name__)
    first = run_audit(make_doc(TS_TEXT), DocPurpose.internal, TS_TEXT.encode())
    second = run_audit(make_doc(SPECIAL_TEXT), DocPurpose.internal, SPECIAL_TEXT.encode())
    for result in (first, second):
        assert not any(e.detector == "model" for e in result.report.entities)
        assert not any(d.startswith("model:") for d in result.report.audit.detectors)
    warnings = [r for r in caplog.records if r.name == md.__name__]
    assert len(warnings) == 1 and "не найден" in warnings[0].getMessage()
    assert MANAGER not in caplog.text and EMPLOYEE not in caplog.text


def test_broken_model_file_skips_step(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, caplog: pytest.LogCaptureFixture,
) -> None:
    (tmp_path / md.MODEL_FILE).write_bytes(b"not an onnx model")
    monkeypatch.setattr(md, "MODEL_DIRS", (tmp_path,))
    md.load_model.cache_clear()
    caplog.set_level(logging.WARNING, logger=md.__name__)
    try:
        assert md.load_model() is None and md.load_model() is None
        assert model_found(TS_TEXT) == []
    finally:
        md.load_model.cache_clear()
    assert len([r for r in caplog.records if r.name == md.__name__]) == 1


# --------------------------------------------- логика на фиктивной модели ---


class FakeModel:
    """Подменяет load_model/predict_proba: вероятности по ключевым словам фрагмента."""

    def __init__(self, rules: dict[str, dict[EntityType, float]]) -> None:
        self.rules = rules
        self.thresholds = {EntityType.TRADE_SECRET: 0.35, EntityType.SPECIAL_CATEGORY: 0.6}
        self.version = "grif-cls-fake0000"

    def predict(self, _model: object, texts: list[str]) -> list[dict[EntityType, float]]:
        out = []
        for text in texts:
            probs = {EntityType.TRADE_SECRET: 0.0, EntityType.SPECIAL_CATEGORY: 0.0}
            for key, value in self.rules.items():
                if key in text:
                    probs.update(value)
            out.append(probs)
        return out


@pytest.fixture
def fake(monkeypatch: pytest.MonkeyPatch) -> FakeModel:
    model = FakeModel({
        "накрутка": {EntityType.TRADE_SECRET: 0.9},
        "диабет": {EntityType.SPECIAL_CATEGORY: 0.8},
    })
    monkeypatch.setattr(md, "load_model", lambda: model)
    monkeypatch.setattr(md, "predict_proba", model.predict)
    return model


def entity(text: str, value: str, etype: EntityType, category: EntityCategory,
           detector: str = "regex") -> Entity:
    start = text.index(value)
    return Entity(type=etype, category=category, start=start, end=start + len(value),
                  value_masked=value, confidence=0.9, detector=detector)


def run_fake(text: str, entities: list[Entity]) -> list[Entity]:
    doc = make_doc(text)
    merged, masked = md.apply_model(doc, entities, mask(text, entities))
    assert mask(text, merged).masked_text == masked.masked_text  # повторная маска не падает
    return merged


def test_financial_under_model_trade_secret_removed(fake: FakeModel) -> None:
    text = "Наша накрутка на кабель составит 120 000 руб. за партию."
    financial = entity(text, "120 000 руб.", EntityType.FINANCIAL, EntityCategory.financial,
                       "dict")
    merged = run_fake(text, [financial])
    assert [(e.type, e.detector) for e in merged] == [(EntityType.TRADE_SECRET, "model")]


def test_model_trade_secret_overlapping_dictionary_one_is_kept(fake: FakeModel) -> None:
    text = "Наша накрутка на кабель составит 14% от закупки."
    secret = entity(text, "накрутка на кабель составит 14%", EntityType.TRADE_SECRET,
                    EntityCategory.trade_secret, "dict")
    merged = run_fake(text, [secret])
    assert merged == [secret]


def test_special_cut_around_person_inside(fake: FakeModel) -> None:
    text = f"По словам врача, работник {EMPLOYEE} давно страдает тяжёлой формой диабета."
    person = entity(text, EMPLOYEE, EntityType.PERSON, EntityCategory.pd, "ner")
    merged = run_fake(text, [person])
    special = [e for e in merged if e.type == EntityType.SPECIAL_CATEGORY]
    assert len(special) == 1
    assert text[special[0].start:special[0].end] == "давно страдает тяжёлой формой диабета."
    assert special[0].value_masked.startswith("д████ с")


def test_special_without_person_dropped(fake: FakeModel) -> None:
    assert run_fake("Закупка тест-полосок для пациентов с диабетом.", []) == []


def test_special_with_person_word_in_neighbor(fake: FakeModel) -> None:
    text = "Кандидат прошёл собеседование. Выявлен диабет второго типа."
    merged = run_fake(text, [])
    assert [e.type for e in merged] == [EntityType.SPECIAL_CATEGORY]


def test_special_dropped_when_too_short_after_cut(fake: FakeModel) -> None:
    text = f"{EMPLOYEE}, диабет."
    person = entity(text, EMPLOYEE, EntityType.PERSON, EntityCategory.pd, "ner")
    assert run_fake(text, [person]) == [person]


# ------------------------------------------------------------ функции ---


def test_carve_edges_and_inner() -> None:
    text = "Петров П. П. перенёс операцию на сердце в мае"
    left = entity(text, "Петров П. П.", EntityType.PERSON, EntityCategory.pd)
    assert md.carve(text, 0, len(text), [left]) == (text.index("перенёс"), len(text))
    assert md.carve(text, 0, len(text), []) == (0, len(text))
    inner_text = "в мае Петров П. П. перенёс операцию на сердце"
    inner = entity(inner_text, "Петров П. П.", EntityType.PERSON, EntityCategory.pd)
    start = inner_text.index("перенёс")
    assert md.carve(inner_text, 0, len(inner_text), [inner]) == (start, len(inner_text))


def test_carve_rejects_short_pieces() -> None:
    text = "у него Петров П. П. ещё"
    person = entity(text, "Петров П. П.", EntityType.PERSON, EntityCategory.pd)
    assert md.carve(text, 0, len(text), [person]) is None


def test_has_person_context() -> None:
    texts = ["Поставка инсулина.", "Ответственный [PERSON_1].", "Срок — май.", "Без людей.",
             "Работник уволен.", "Перенёс инфаркт."]
    assert has_person_context(texts, word_in_neighbors=False) == [
        True, True, True, False, True, False]
    assert has_person_context(texts, word_in_neighbors=True) == [
        True, True, True, True, True, True]
    assert has_person_context(["Номер [SNILS_1]."], word_in_neighbors=False) == [True]
    assert has_person_context(["Телефон [PHONE_1]."], word_in_neighbors=False) == [False]
    assert has_person_context([], word_in_neighbors=True) == []


def test_split_fragments_offsets_from_backend_module() -> None:
    text = "Наценка составит 14%. Срок поставки — 10 дней."
    assert [text[s:e] for s, e in split_fragments(text)] == [
        "Наценка составит 14%.", "Срок поставки — 10 дней."]


def test_add_model_entities_sorted_and_keeps_other_financial() -> None:
    text = "Сумма 500 руб. Наша накрутка 20%."
    money = entity(text, "500 руб.", EntityType.FINANCIAL, EntityCategory.financial, "dict")
    secret = entity(text, "Наша накрутка 20%.", EntityType.TRADE_SECRET,
                    EntityCategory.trade_secret, "model")
    assert md.add_model_entities([money], [secret]) == [money, secret]


# ------------------------------------- T-019b: привязка к человеку в абзаце ---

CHARACTERISTIC = (f"Характеристика: {EMPLOYEE} работает в отделе снабжения с 2019 года. "
                  "Отличается аккуратностью и пунктуальностью. "
                  "В прошлом месяце перенёс операцию на сердце.")


def make_paragraphs(paragraphs: list[str], pages: int | None = None) -> ExtractedDocument:
    """Документ из абзацев-сегментов (как docx; pages=1 — как PDF)."""
    segments: list[TextSegment] = []
    offset = 0
    for i, text in enumerate(paragraphs):
        segments.append(TextSegment(id=i, text=text, offset=offset, page=1 if pages else None))
        offset += len(text) + 1
    return ExtractedDocument(document_id="synthetic-paragraph", filename="synthetic.docx",
                             text="\n".join(paragraphs), segments=segments, pages=pages)


def special_found(doc: ExtractedDocument) -> list[str]:
    return [doc.text[e.start:e.end] for e in detect_with_model(doc)
            if e.type == EntityType.SPECIAL_CATEGORY and e.detector == "model"]


@requires_model
def test_person_two_sentences_above_in_same_paragraph() -> None:
    """Человек назван в начале абзаца, сведения — через предложение: тот же субъект."""
    assert special_found(make_paragraphs([CHARACTERISTIC])) == [
        "В прошлом месяце перенёс операцию на сердце."]


@requires_model
def test_paragraph_rule_off_loses_distant_person(monkeypatch: pytest.MonkeyPatch) -> None:
    """Без абзацев остаётся только окно ±1 фрагмент — та же фраза не засчитывается."""
    monkeypatch.setattr(md, "paragraph_bounds", lambda doc: [])
    assert special_found(make_paragraphs([CHARACTERISTIC])) == []


@requires_model
def test_line_is_paragraph_in_plain_text_but_not_in_pdf() -> None:
    """В .txt и docx строка — отдельный абзац; в PDF перевод строки — перенос внутри блока."""
    lines = CHARACTERISTIC.replace(". ", ".\n")
    assert special_found(make_paragraphs([lines])) == []
    assert special_found(make_paragraphs([lines], pages=1)) == [
        "В прошлом месяце перенёс операцию на сердце."]


@requires_model
def test_other_paragraph_person_does_not_bind_medical_phrase() -> None:
    """Человек в соседнем, но не смежном абзаце — медицинская закупка остаётся не ПДн."""
    doc = make_paragraphs([f"Ответственный — {EMPLOYEE}.", "Срок поставки — май.",
                           "Склад принимает товар по будням.", MEDICAL_TEXT])
    assert special_found(doc) == []
