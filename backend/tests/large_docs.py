"""Детерминированные синтетические документы для T-031, без сети."""

from html import escape
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook

DOCX_SECTIONS = 70
PDF_PAGES = 60
ESTIMATE_ROWS = 5000
# Вымышленные учебные представители; совпадения с реальными людьми случайны.
DOCX_NAME = "Примеров Мирон Платонович"
PDF_NAME = "Образцов Савелий Родионович"
PHONE = "+7 (999) 000-12-34"
EMAIL = "representative@example.ru"
PASSPORT = "99 17 170303"
SNILS = "112-233-445 95"
BIRTH_DATE = "17.03.1987"
SUPPLY_PARAGRAPH = (
    "Поставщик передаёт товар в согласованной комплектации. Заказчик проверяет "
    "количество и качество при приёмке. Стороны оформляют акт и согласуют "
    "устранение замечаний в течение пяти рабочих дней."
)


def build_docx(path: Path) -> None:
    """70 разделов с явными разрывами страниц, по 12 абзацев в каждом."""
    if not path.parent.is_dir():
        raise FileNotFoundError(path.parent)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.header.paragraphs[0].text = "Конфиденциально"
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10)
    style.paragraph_format.space_after = Pt(4)
    for number in range(1, DOCX_SECTIONS + 1):
        heading = document.add_heading(f"{number}. Договор поставки: условия исполнения", 1)
        heading.paragraph_format.page_break_before = number > 1
        for paragraph in range(12):
            text = f"{number}.{paragraph + 1}. {SUPPLY_PARAGRAPH}"
            if paragraph == 0 and number % 10 == 0:
                text = f"Представитель поставщика: {DOCX_NAME}. Телефон: {PHONE}. {EMAIL}."
            elif paragraph == 1 and number % 15 == 0:
                text = "Себестоимость единицы 48 300 руб. Условия расчёта поставщика."
            document.add_paragraph(text)
    document.core_properties.author = "Синтетический корпус T-031"
    document.save(path)


def build_pdf(path: Path) -> None:
    """60 страниц HTML/CSS; MuPDF встраивает шрифт с кириллицей."""
    if not path.parent.is_dir():
        raise FileNotFoundError(path.parent)
    css = "body {font-family:sans-serif;font-size:11pt;line-height:1.35;}"
    with pymupdf.open() as document:
        for number in range(1, PDF_PAGES + 1):
            texts = [f"Заявка на поставку оборудования. Страница {number}.", SUPPLY_PARAGRAPH]
            if number % 5 == 0:
                texts.extend([
                    f"Представитель: {PDF_NAME}.", f"Паспорт: {PASSPORT}.",
                    f"СНИЛС: {SNILS}.", f"Дата рождения: {BIRTH_DATE}.",
                ])
            html = "".join(f"<p>{escape(text)}</p>" for text in texts)
            page = document.new_page(width=595, height=842)
            spare, _ = page.insert_htmlbox(
                pymupdf.Rect(55, 45, 540, 795), html, css=css, scale_low=1,
            )
            if spare < 0:
                raise ValueError("Текст не помещается на страницу PDF")
        document.set_metadata({"author": "Синтетический корпус T-031"})
        document.save(path)


def build_xlsx(path: Path) -> None:
    """5000 строк позиций плюс заголовок и показатели коммерческой тайны."""
    if not path.parent.is_dir():
        raise FileNotFoundError(path.parent)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Смета"
    sheet.append(["Себестоимость единицы", 48300, "руб."])
    sheet.append(["Наценка", 0.25])
    sheet["B2"].number_format = "0%"
    sheet.append(["Наименование", "Количество", "Цена", "Сумма"])
    for number in range(1, ESTIMATE_ROWS + 1):
        quantity, price = 1 + number % 10, 100 + number % 900
        sheet.append([f"Комплект оборудования № {number}", quantity, price, quantity * price])
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 42
    workbook.save(path)
    workbook.close()
